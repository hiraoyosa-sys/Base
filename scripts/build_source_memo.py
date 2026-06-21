# -*- coding: utf-8 -*-
"""出典メモ（平尾用・配布しない）：資料の各記述が何に基づくかを平尾に伝えるための内部メモ。"""
import os
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

OUT = os.path.join(os.path.dirname(__file__), "..", "outputs", "出典メモ_平尾用.docx")
JP = "游ゴシック"
RED = RGBColor(0xC0, 0x00, 0x00); NAVY = RGBColor(0x00, 0x3F, 0x7E); GRAY = RGBColor(0x55, 0x55, 0x55)

doc = Document()
cp = doc.core_properties; cp.author = ""; cp.last_modified_by = ""
st = doc.styles["Normal"]; st.font.name = JP; st.font.size = Pt(10.5)
st.element.rPr.rFonts.set(qn("w:eastAsia"), JP)

def sf(run, size=10.5, bold=False, color=None):
    run.font.name = JP; run.font.size = Pt(size); run.font.bold = bold
    run._element.rPr.rFonts.set(qn("w:eastAsia"), JP)
    if color is not None: run.font.color.rgb = color

p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
sf(p.add_run("出典メモ（平尾用・配布しない）"), 15, True, NAVY)
p = doc.add_paragraph()
sf(p.add_run("配布資料（PPTX/Word/Excel）には出典を書いていません。各記述の根拠を下表に整理しました。社外資料の扱いに注意。"), 10, False, GRAY)

rows = [
    ("原料EO中A-ALD/EO=0.78kg/t（過去上位）", "運転実績。2026年2〜5月EG系まとめデッキ「EGプラントへの不純物持込み」。", ""),
    ("ラボでEG反応器出口を模擬しALD添加→UV悪化", "黒木RDのラボ再現（6/10・6/17品質定例）。＝EG反応器側の検証で、EO反応器ではない。", "平尾指摘を反映"),
    ("水分が多いほど縮合は進みにくい・感度2〜3割", "6/17品質定例の基礎研の見立て（DEG色相悪化＝脱水反応／水分感度2〜3割）。", "★クリーンな“水分を増やす”ベンチ実験ではない＝“見立て”として記載"),
    ("付着物+DEG→380nmのみ・450nm再現せず", "6/16〜17のDEG共存保管試験（基礎研）。蒸留DEG短時間でも450nm出ず。", ""),
    ("着色物質=分子量200〜800・457nm・ppbのポリエナール", "6/16基礎研GPC、6/17“約10ppb”の試算（Beer則）。", ""),
    ("酸素があると有機酸へ酸化され縮合が進まない", "6/17品質定例の基礎研（酸素で前駆体が酸化され重合が進まない）。", ""),
    ("C-1502で収支の入量を超える/温度ダウンでA-ALD低下", "ALDバランス整理（6/11 EO対策会議）。塔間循環の解釈は平尾＋伊藤良太(RD)仮説。", ""),
    ("前駆体どうしの脱水縮合で共役伸長→450nm（タンク）", "6/16議事録の推定メカニズム＋基礎研「重合物同士で成長、使い切るとAPHA30で頭打ち」。", ""),
    ("温度を下げない（前駆体が色相に出る）", "6/18〜19の壁打ち（温度を下げると分解せず前駆体として下流へ）。", ""),
    ("リン酸の効き（pH低下／金属不活性化）", "6/16〜17の二段解釈。C-1503壁面のP分析で切り分け（要確認）。", ""),
    ("N分含有の示唆", "6/18基礎研B資料（N分含有あり、過去のアミン由来とは分子量が異なる）。", "残論点"),
    ("UV450nm 悪化品0.004/日・正常品0.002不動", "6/19 DEG対策会議 向後（T-604実績）。", ""),
]
# 外したもの
removed = [
    ("（タンク④に入れていた）アルドール縮合でポリ不飽和体・pH制御で抑制", "Dow特許 US8293949（社外）。", "★pHの話が唐突になるため資料からは外した"),
    ("シェル：F-ALDは出口温度にアレニウス型／滞留短縮・錆除去で低減", "Shell GEOTEM502系資料。", "★今回の着色の主筋に非中核＝資料から外した"),
    ("シェル：DEG/TEGのエーテルがペルオキシ酸化で切断されアルデヒド生成", "野田「可能性2（文献裏付けあり）」だが本人が「主因には量的根拠不足」。", "★アルデヒドの別生成源の説＝主筋でないため資料から外した"),
]

def shade(c, hexcolor):
    tcPr = c._tc.get_or_add_tcPr(); sh = OxmlElement("w:shd"); sh.set(qn("w:val"), "clear"); sh.set(qn("w:fill"), hexcolor); tcPr.append(sh)

def add_table(title, data, headfill):
    h = doc.add_paragraph(); sf(h.add_run(title), 12, True, NAVY)
    t = doc.add_table(rows=1, cols=3); t.alignment = WD_TABLE_ALIGNMENT.CENTER; t.style = "Table Grid"
    for j, ht in enumerate(["資料での記述", "出典（社内・日付）", "備考"]):
        c = t.rows[0].cells[j]; c.text = ""; r = c.paragraphs[0].add_run(); r.text = ht; sf(r, 9.5, True, NAVY); shade(c, headfill)
    for a, b, note in data:
        cells = t.add_row().cells
        for j, v in enumerate([a, b, note]):
            cells[j].text = ""; r = cells[j].paragraphs[0].add_run(); r.text = v
            sf(r, 9, False, (RED if (j == 2 and v.startswith("★")) else None))
    from docx.shared import Cm
    for row in t.rows:
        row.cells[0].width = Cm(6.5); row.cells[1].width = Cm(7.5); row.cells[2].width = Cm(5.5)
    doc.add_paragraph()

add_table("1. 資料に載せた記述の根拠", rows, "EBEFF2")
add_table("2. 検討したが資料から外したもの（参考）", removed, "FCE4D6")

os.makedirs(os.path.dirname(OUT), exist_ok=True)
doc.save(OUT)
print("saved:", os.path.normpath(OUT))
