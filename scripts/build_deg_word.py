# -*- coding: utf-8 -*-
"""解説用Word：製品DEG色相悪化 メカニズム解説（事実→仮説→対応）。配布クリーン。"""
import os
from docx import Document
from docx.shared import Pt, RGBColor, Cm, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import deg_master as M

HERE = os.path.dirname(__file__)
OUTDIR = os.path.join(HERE, "..", "outputs")
OUT = os.path.join(OUTDIR, "製品DEG色相悪化_メカニズム解説.docx")
JP = "游ゴシック"
# 薄いMC配色（見出しは中間ブルー、濃色は使わない）
NAVY = RGBColor(0x00, 0x5B, 0xAB)
DBLUE = RGBColor(0x00, 0x3F, 0x7E)
GRAY = RGBColor(0x55, 0x55, 0x55)
FILL_HEAD = "EBEFF2"; FILL_PROC = "D7E0E5"; FILL_OPE = "E8F1E5"; FILL_RD = "FBF1DC"; FILL_LIT = "E9F0F7"; FILL_SUJI = "DCE6EF"

doc = Document()
cp = doc.core_properties
cp.author = ""; cp.last_modified_by = ""; cp.title = "製品DEG色相悪化 メカニズム解説"; cp.comments = ""

# 既定フォント
style = doc.styles["Normal"]
style.font.name = JP; style.font.size = Pt(10.5)
style.element.rPr.rFonts.set(qn("w:eastAsia"), JP)

def set_font(run, size=10.5, bold=False, color=None):
    run.font.name = JP; run.font.size = Pt(size); run.font.bold = bold
    run._element.rPr.rFonts.set(qn("w:eastAsia"), JP)
    if color is not None: run.font.color.rgb = color

def title(text):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(p.add_run(text), 16, True, NAVY)

def sub(text):
    set_font(p_caption(text), 11, True)

def h(text):
    p = doc.add_paragraph(); p.space_before = Pt(6)
    set_font(p.add_run(text), 13, True, NAVY)

def body(text, bullet=False, size=10.5):
    p = doc.add_paragraph(style="List Bullet" if bullet else None)
    set_font(p.add_run(text), size)
    p.paragraph_format.space_after = Pt(3)
    return p

def note(text):
    p = doc.add_paragraph(); set_font(p.add_run(text), 9.5, False, GRAY)
    p.paragraph_format.space_after = Pt(4)

# ── 表紙的ヘッダ ──
title("製品DEG色相悪化　メカニズム解説")
pc = doc.add_paragraph(); pc.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_font(pc.add_run("事実の整理 → 推定メカニズム → 対応"), 11, False, GRAY)
doc.add_paragraph()

# 1
h("1. 発生事象（事実）")
for t in M.PHENOMENON: body(t, bullet=True)

# 2
h("2. 着色物質と着色現象")
p = doc.add_paragraph(); set_font(p.add_run("■ 確認した事実"), 11, True, NAVY)
for t in M.COLORANT_FACT: body(t, bullet=True)
p = doc.add_paragraph(); set_font(p.add_run("■ ここから読める推定"), 11, True, GRAY)
for t in M.COLORANT_EST: body(t, bullet=True)

# 3
h("3. " + M.MECH_CAPTION)
note("各段は「運転で見えた事実・基礎研の確認・反応原理」を重ねて導いた最もらしい筋。証明ではないが、観測・分析・原理と整合する。")
for head, txt in M.MECHANISM:
    p = doc.add_paragraph(); set_font(p.add_run(head), 11.5, True, NAVY)
    p.paragraph_format.space_before = Pt(4); p.paragraph_format.space_after = Pt(1)
    body(txt)
p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(4)
set_font(p.add_run("一筋でいえば："), 11, True)
set_font(p.add_run(M.MECH_ONELINE), 11)
note("※ " + M.MECH_CAVEAT)
# 反応スキーム図（化学式）
_aldol = os.path.join(OUTDIR, "aldol_scheme.png")
if os.path.exists(_aldol):
    pp = doc.add_paragraph(); pp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pp.add_run().add_picture(_aldol, width=Inches(6.6))
    cap = doc.add_paragraph(); cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(cap.add_run("図1　アルドール縮合で共役が伸び450nm（黄）に至る反応スキーム"), 9.5, False, GRAY)

# 工程別マトリクス
def shade(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    sh = OxmlElement("w:shd"); sh.set(qn("w:val"), "clear"); sh.set(qn("w:fill"), hexcolor)
    tcPr.append(sh)

def cell_text(cell, text, bold=False, color=None, size=9):
    cell.text = ""
    p = cell.paragraphs[0]; r = p.add_run(); r.text = text
    set_font(r, size, bold, color); p.paragraph_format.space_after = Pt(0)

# 工程別グリッド（横＝工程）はランドスケープのセクションに置く
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.shared import Cm as _Cm
sec = doc.add_section(WD_SECTION.NEW_PAGE)
sec.orientation = WD_ORIENT.LANDSCAPE
sec.page_width, sec.page_height = sec.page_height, sec.page_width
sec.left_margin = _Cm(1.2); sec.right_margin = _Cm(1.2)

h("3.5　工程別グリッド（横＝工程／縦＝項目）")
note("空欄＝確実な情報なし（無理に埋めない）。行の色＝根拠：緑＝事実（運転・解放・RD）／青＝文献・社外知見／灰＝一般原理／青灰＝小結論。空欄の多い工程ほど不確実。")
KIND_FILL = {"fact": FILL_OPE, "lit": FILL_LIT, "principle": "F1F1F1", "concl": FILL_SUJI}
ncol = 1 + len(M.STAGES)
tbl = doc.add_table(rows=1, cols=ncol); tbl.alignment = WD_TABLE_ALIGNMENT.CENTER; tbl.style = "Table Grid"
cell_text(tbl.rows[0].cells[0], "項目", bold=True, color=DBLUE, size=8.5); shade(tbl.rows[0].cells[0], FILL_HEAD)
for j, stg in enumerate(M.STAGES, 1):
    cell_text(tbl.rows[0].cells[j], stg.replace("\n", " "), bold=True, color=DBLUE, size=8); shade(tbl.rows[0].cells[j], FILL_PROC)
# 裏付け強度
srow = tbl.add_row().cells
cell_text(srow[0], "裏付け強度", bold=True, color=DBLUE, size=8); shade(srow[0], FILL_HEAD)
for j, sym in enumerate(M.STRENGTH, 1):
    cell_text(srow[j], sym, bold=True, color=DBLUE, size=11); shade(srow[j], FILL_PROC)
    srow[j].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
for label, kind, cells in M.GRID:
    row = tbl.add_row().cells
    cell_text(row[0], label, bold=True, color=DBLUE, size=8); shade(row[0], FILL_HEAD)
    for j, val in enumerate(cells, 1):
        cell_text(row[j], val, size=7.5); shade(row[j], KIND_FILL[kind])
widths = [2.6] + [3.5] * len(M.STAGES)
for row in tbl.rows:
    for j, w in enumerate(widths):
        row.cells[j].width = _Cm(w)
# 1本の筋（表の外）
ps = doc.add_paragraph(); ps.paragraph_format.space_before = Pt(6)
set_font(ps.add_run("全工程をトータル＝1本の筋（仮説）： "), 10.5, True, NAVY)
set_font(ps.add_run(M.SUJI_TEXT), 10)
# 縦（ポートレート）に戻す
sec2 = doc.add_section(WD_SECTION.NEW_PAGE)
sec2.orientation = WD_ORIENT.PORTRAIT
sec2.page_width, sec2.page_height = sec2.page_height, sec2.page_width

# 4
h("4. なぜ今回だけDEGまで到達し着色したか")
note("過去にもアルデヒドが高い時期はあったが今回のような色相悪化は起きていない。アルデヒド単独でなく、反応させてしまう条件が今回そろった。")
for head, txt, act in M.WHY_NOW:
    body("【" + head + "】 " + txt + "　" + act, bullet=True)
note("※ " + M.WATER_NOTE)

# 5
h("5. 設備（鉄サビ・付着物）の関与の評価")
for t in M.RUST: body(t, bullet=True)

# 6
h("6. 開放・検査の整理（仮説駆動）")
body(M.INSPECTION_LOGIC)
for eq, txt in M.INSPECTION:
    body("【" + eq + "】 " + txt, bullet=True)

# 7
h("7. 対応（再発防止・スタートアップ運転条件）")
for head, txt, conf in M.COUNTERMEASURE:
    body("【" + head + "】 " + txt + "（" + conf + "）", bullet=True)

# 8
h("8. 出荷再開に向けた品質管理・早期判定")
for t in M.QC: body(t, bullet=True)
p = doc.add_paragraph(); set_font(p.add_run("■ 出荷判定の運用（歯止め）"), 11, True, NAVY)
note(M.GATE_NOTE)
for k, v in M.GATE: body("【" + k + "】 " + v, bullet=True)

# 9
h("9. 残論点（要確認）")
for head, txt in M.OPEN_ISSUES:
    body("【" + head + "】 " + txt, bullet=True)

os.makedirs(os.path.dirname(OUT), exist_ok=True)
doc.save(OUT)
print("saved:", os.path.normpath(OUT))
