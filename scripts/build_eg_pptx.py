# -*- coding: utf-8 -*-
"""EG品質／DEG色相 説明 PowerPoint生成。

ルール準拠：MCC標準ワイド 33.867×19.05cm／白背景・文字黒基調・モノクロ／1スライド=1メッセージ／
本文に「→」記号書きを使わない（工程の流れは矢印図形で表す）／程度語・造語を使わない／
外部文献・AI痕跡を入れない／作成者メタは平尾名義。構成は 事実→仮説→対応。
"""
import os
from pptx import Presentation
from pptx.util import Cm, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
import eg_content as C

OUT = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "outputs")
os.makedirs(OUT, exist_ok=True)

FONT = "Meiryo"
INK = RGBColor(0x00, 0x00, 0x00)
SUB = RGBColor(0x40, 0x40, 0x40)
HEAD = RGBColor(0xD9, 0xD9, 0xD9)
ALT = RGBColor(0xF2, 0xF2, 0xF2)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LINE = RGBColor(0x80, 0x80, 0x80)
RULE = RGBColor(0x00, 0x00, 0x00)

prs = Presentation()
prs.slide_width = Cm(33.867)
prs.slide_height = Cm(19.05)
SW, SH = prs.slide_width, prs.slide_height
BLANK = prs.slide_layouts[6]

ML = Cm(1.4)          # 左右マージン
MT = Cm(1.1)          # 上マージン
CW = SW - ML * 2      # コンテンツ幅


def slide():
    s = prs.slides.add_slide(BLANK)
    # 背景を白で固定
    bg = s.background
    bg.fill.solid()
    bg.fill.fore_color.rgb = WHITE
    return s


def _set_font(run, size, bold=False, color=INK):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    # 日本語フォントを明示
    rPr = run._r.get_or_add_rPr()
    ea = rPr.find(qn('a:ea'))
    if ea is None:
        ea = rPr.makeelement(qn('a:ea'), {})
        rPr.append(ea)
    ea.set('typeface', FONT)


def textbox(s, l, t, w, h, lines, size=16, bold=False, color=INK, align=PP_ALIGN.LEFT,
            anchor=MSO_ANCHOR.TOP, leading=1.12, bullet=False, space_after=4):
    tb = s.shapes.add_textbox(l, t, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = Cm(0.1)
    tf.margin_right = Cm(0.1)
    tf.margin_top = Cm(0.05)
    tf.margin_bottom = Cm(0.05)
    if isinstance(lines, str):
        lines = [lines]
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = leading
        p.space_after = Pt(space_after)
        p.space_before = Pt(0)
        txt = ("・" + ln) if bullet else ln
        r = p.add_run()
        r.text = txt
        _set_font(r, size, bold, color)
    return tb


def title_bar(s, text, sub=None):
    textbox(s, ML, MT, CW, Cm(1.4), text, size=24, bold=True, color=INK)
    # タイトル下の細い罫線
    ln = s.shapes.add_connector(2, ML, MT + Cm(1.45), ML + CW, MT + Cm(1.45))
    ln.line.color.rgb = RULE
    ln.line.width = Pt(1.5)
    if sub:
        textbox(s, ML, MT + Cm(1.5), CW, Cm(0.7), sub, size=12, color=SUB)
    return MT + Cm(2.25 if sub else 1.75)


def page_note(s):
    textbox(s, ML, SH - Cm(0.9), CW, Cm(0.6),
            f"プロセス技術部　EOGモデルチェンジPJ　／　{C.DATE} 時点",
            size=9, color=SUB, align=PP_ALIGN.RIGHT)


def set_cell_border(cell, color="808080", w=6350):
    """python-pptxは罫線APIが無いのでXMLで4辺の細罫線を付ける（モノクロ）。
    OOXMLのスキーマ順では罫線(a:lnL/R/T/B)は塗り(a:solidFill)より前に置く必要があるため、
    先頭へ逆順で挿入して L,R,T,B の順になるようにする。"""
    tcPr = cell._tc.get_or_add_tcPr()
    for tag in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
        for ex in tcPr.findall(qn(tag)):
            tcPr.remove(ex)
    for tag in ("a:lnB", "a:lnT", "a:lnR", "a:lnL"):
        ln = tcPr.makeelement(qn(tag), {"w": str(w), "cap": "flat"})
        fill = ln.makeelement(qn('a:solidFill'), {})
        clr = ln.makeelement(qn('a:srgbClr'), {"val": color})
        fill.append(clr)
        ln.append(fill)
        tcPr.insert(0, ln)


def table_slide(title, headers, rows, col_w, font_size=11, head_size=11,
                aligns=None, sub=None, note=None, strong_first=False):
    s = slide()
    top = title_bar(s, title, sub)
    aligns = aligns or [PP_ALIGN.LEFT] * len(headers)
    nrow = len(rows) + 1
    ncol = len(headers)
    total_w = sum(col_w)
    avail_h = SH - top - Cm(1.2)
    gtab = s.shapes.add_table(nrow, ncol, ML, top, Emu(int(total_w)), Cm(1.0)).table
    # 既定スタイル（色付き）を抑止
    tbl = gtab._tbl
    pr = tbl.tblPr
    pr.set('firstRow', '0')
    pr.set('bandRow', '0')
    for i, w in enumerate(col_w):
        gtab.columns[i].width = Emu(int(w))
    # ヘッダ
    for c, h in enumerate(headers):
        cell = gtab.cell(0, c)
        cell.fill.solid()
        cell.fill.fore_color.rgb = HEAD
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        cell.margin_left = Cm(0.12); cell.margin_right = Cm(0.12)
        cell.margin_top = Cm(0.04); cell.margin_bottom = Cm(0.04)
        tf = cell.text_frame; tf.word_wrap = True
        p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
        r = p.add_run(); r.text = h; _set_font(r, head_size, True, INK)
        set_cell_border(cell)
    # 本文
    for ri, row in enumerate(rows, 1):
        for c, v in enumerate(row):
            cell = gtab.cell(ri, c)
            cell.fill.solid()
            cell.fill.fore_color.rgb = ALT if ri % 2 == 0 else WHITE
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.margin_left = Cm(0.12); cell.margin_right = Cm(0.12)
            cell.margin_top = Cm(0.03); cell.margin_bottom = Cm(0.03)
            tf = cell.text_frame; tf.word_wrap = True
            p = tf.paragraphs[0]; p.alignment = aligns[c]; p.line_spacing = 1.05
            r = p.add_run(); r.text = str(v)
            _set_font(r, font_size, bold=(strong_first and c == 0), color=INK)
            set_cell_border(cell)
    if note:
        textbox(s, ML, SH - Cm(1.7), CW, Cm(0.9), note, size=10, bold=True, color=INK)
    page_note(s)
    return s


def bullets_slide(title, bullets, sub=None, size=16, two_col=False):
    s = slide()
    top = title_bar(s, title, sub)
    if two_col and len(bullets) > 4:
        half = (len(bullets) + 1) // 2
        textbox(s, ML, top, CW / 2 - Cm(0.3), SH - top - Cm(1.2), bullets[:half],
                size=size, bullet=True, leading=1.18, space_after=8)
        textbox(s, ML + CW / 2 + Cm(0.3), top, CW / 2 - Cm(0.3), SH - top - Cm(1.2), bullets[half:],
                size=size, bullet=True, leading=1.18, space_after=8)
    else:
        textbox(s, ML, top, CW, SH - top - Cm(1.2), bullets,
                size=size, bullet=True, leading=1.2, space_after=10)
    page_note(s)
    return s


# ============ 1. 表紙 ============
s = slide()
textbox(s, ML, Cm(6.2), CW, Cm(3.2), C.TITLE, size=34, bold=True, color=INK, align=PP_ALIGN.CENTER)
textbox(s, ML, Cm(9.6), CW, Cm(1.2), C.SUBTITLE, size=18, color=SUB, align=PP_ALIGN.CENTER)
textbox(s, ML, Cm(11.2), CW, Cm(1.0), f"{C.DATE}", size=14, color=SUB, align=PP_ALIGN.CENTER)
textbox(s, ML, SH - Cm(1.6), CW, Cm(0.8), "プロセス技術部　EOGモデルチェンジPJ", size=12, color=SUB, align=PP_ALIGN.CENTER)

# ============ 2. 目的 ============
bullets_slide("この資料の目的", C.PURPOSE, size=16)

# ============ 3. 結論サマリ ============
bullets_slide("結論（先に要点）", C.SUMMARY, size=15)

# ============ 4. 現状：製品別の品質状況 ============
table_slide("現状（事実）：製品別の品質状況",
            ["製品／系", "状況", "特徴的な事象"],
            [[a, b, c2] for a, b, c2 in C.PRODUCT_STATUS],
            [Cm(4.5), Cm(13.0), Cm(13.5)],
            font_size=11, aligns=[PP_ALIGN.CENTER, PP_ALIGN.LEFT, PP_ALIGN.LEFT],
            sub="顧客の窓口は、MEGが東レ、TEGが東洋紡、DEG色相がレゾナック。",
            strong_first=True)

# ============ 5. 現状：運転対応の実績 ============
table_slide("現状（事実）：運転対応の実績",
            ["操作", "期間", "HUV220低減", "DEG色相改善"],
            [list(r) for r in C.UNTEN_JISSEKI],
            [Cm(14.0), Cm(6.0), Cm(5.5), Cm(5.5)],
            font_size=10.5, head_size=10.5,
            aligns=[PP_ALIGN.LEFT, PP_ALIGN.CENTER, PP_ALIGN.CENTER, PP_ALIGN.CENTER],
            note="要点：" + C.UNTEN_NOTE)

# ============ 6. 確認された事実（要点） ============
key_facts = ["F1", "F2", "F7", "F8", "F11", "F16", "F17", "F18"]
fmap = dict(C.FACTS)
bullets_slide("確認された事実（要点）",
              [fmap[k] for k in key_facts], size=14)

# ============ 7. 着色の見立て：工程の流れ ============
s = slide()
top = title_bar(s, "着色の見立て：工程に沿った筋")
n = len(C.SCENARIO)
gap = Cm(0.45)
bw = (CW - gap * (n - 1)) / n
bh = Cm(8.4)
by = top + Cm(0.2)
for i, (name, what, _why) in enumerate(C.SCENARIO):
    bx = ML + i * (bw + gap)
    box = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, bx, by, bw, bh)
    box.fill.solid(); box.fill.fore_color.rgb = WHITE
    box.line.color.rgb = INK; box.line.width = Pt(1.0)
    box.shadow.inherit = False
    tf = box.text_frame; tf.word_wrap = True
    tf.margin_left = Cm(0.12); tf.margin_right = Cm(0.12); tf.margin_top = Cm(0.1)
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = name; _set_font(r, 11, True, INK)
    p2 = tf.add_paragraph(); p2.alignment = PP_ALIGN.LEFT; p2.line_spacing = 1.05; p2.space_before = Pt(6)
    r2 = p2.add_run(); r2.text = what; _set_font(r2, 9.5, False, INK)
    # 工程間の矢印（図形）
    if i < n - 1:
        ax = bx + bw + Cm(0.02)
        ar = s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, ax, by + bh / 2 - Cm(0.25), gap - Cm(0.04), Cm(0.5))
        ar.fill.solid(); ar.fill.fore_color.rgb = SUB
        ar.line.fill.background()
        ar.shadow.inherit = False
textbox(s, ML, by + bh + Cm(0.3), CW, Cm(2.4), "筋：" + C.SCENARIO_SUJI, size=12, color=INK, leading=1.2)
page_note(s)

# ============ 8. 機構の重要更新 ============
table_slide("着色の見立て：重要な更新（6/17〜6/20）",
            ["論点", "内容"],
            [[a, b] for a, b in C.UPDATES],
            [Cm(7.0), Cm(24.0)],
            font_size=10.5, aligns=[PP_ALIGN.LEFT, PP_ALIGN.LEFT], strong_first=True)

# ============ 9. 薬剤：NaOHとリン酸 ============
drug_rows = [
    ["NaOH", C.NAOH["投入箇所"], C.NAOH["効果"], C.NAOH["機構"]],
    ["リン酸Na", C.PHOS["投入箇所"] + " 投入量 " + C.PHOS["投入量"], C.PHOS["効果"],
     "DEG側：" + C.PHOS["機構DEG"] + " MEG側：" + C.PHOS["機構MEG"]],
]
table_slide("対応：薬剤（NaOHとリン酸Na）の整理",
            ["薬剤", "投入箇所", "効果", "機構"],
            drug_rows, [Cm(3.2), Cm(8.5), Cm(9.5), Cm(9.8)],
            font_size=10, aligns=[PP_ALIGN.CENTER, PP_ALIGN.LEFT, PP_ALIGN.LEFT, PP_ALIGN.LEFT],
            sub="DEG色相に効いたのはリン酸Naのみ。継続投入が前提（止めると戻る）。",
            strong_first=True)

# ============ 10. 薬剤：リン酸テスト履歴 ============
table_slide("対応：リン酸Na テスト履歴",
            ["回", "時期", "条件", "結果"],
            [list(r) for r in C.PHOS_TEST],
            [Cm(2.0), Cm(5.0), Cm(10.0), Cm(14.0)],
            font_size=10.5, aligns=[PP_ALIGN.CENTER, PP_ALIGN.LEFT, PP_ALIGN.LEFT, PP_ALIGN.LEFT])

# ============ 11. 対応方案 ============
table_slide("対応：方案（方案1〜7）",
            ["方案", "内容", "戦略", "現状"],
            [[a, b, c2, e] for (a, b, c2, _d, e) in C.HOUAN],
            [Cm(3.0), Cm(10.0), Cm(6.0), Cm(12.0)],
            font_size=10.5, aligns=[PP_ALIGN.CENTER, PP_ALIGN.LEFT, PP_ALIGN.CENTER, PP_ALIGN.LEFT],
            note="方針：" + C.HOUAN_HOSHIN, strong_first=True)

# ============ 12. SU運転条件案 ============
table_slide("対応：スタートアップ運転条件案",
            ["項目", "方針"],
            [[a, b] for a, b in C.SU_JOKEN],
            [Cm(7.0), Cm(24.0)],
            font_size=12, aligns=[PP_ALIGN.LEFT, PP_ALIGN.LEFT],
            sub="6/19 DEG対策会議。6/25は承認の場、6/22〜24で根拠を付けて合意。",
            note=C.SU_NOTE, strong_first=True)

# ============ 13. 早期判断指標 ============
table_slide("対応：早期色相判断・指標",
            ["指標", "内容"],
            [[a, b] for a, b in C.SHIHYO],
            [Cm(6.0), Cm(25.0)],
            font_size=12, aligns=[PP_ALIGN.LEFT, PP_ALIGN.LEFT], strong_first=True)

# ============ 14. 異物・顧客対応 ============
s = slide()
top = title_bar(s, "対応：異物の確認と顧客への説明")
half = CW / 2 - Cm(0.4)
textbox(s, ML, top, half, Cm(1.0), "異物・タンク・ライン処理", size=14, bold=True)
textbox(s, ML, top + Cm(1.0), half, SH - top - Cm(2.2),
        [f"{a}：{b}" for a, b in C.IBUTSU], size=12, bullet=True, leading=1.18, space_after=8)
textbox(s, ML + CW / 2 + Cm(0.4), top, half, Cm(1.0), "顧客（レゾナック）向け説明の骨子", size=14, bold=True)
textbox(s, ML + CW / 2 + Cm(0.4), top + Cm(1.0), half, SH - top - Cm(2.6),
        C.KOKYAKU, size=12, bullet=True, leading=1.18, space_after=8)
textbox(s, ML + CW / 2 + Cm(0.4), SH - Cm(1.7), half, Cm(0.9), C.KOKYAKU_NOTE, size=10, color=SUB)
page_note(s)

# ============ 15. M50との接続 ============
bullets_slide("M50（モデルチェンジ）との接続", C.M50, size=16)

# ============ 16. 次アクション ============
bullets_slide("次アクション", C.NEXT, size=14)

# ---- メタデータ（作成者：平尾名義／AI痕跡なし）----
cp = prs.core_properties
cp.author = C.AUTHOR
cp.last_modified_by = C.AUTHOR
cp.title = "EG品質の現状とDEG色相悪化への対応"

path = os.path.join(OUT, "EG品質_DEG色相_説明.pptx")
prs.save(path)
print("saved:", path, "slides:", len(prs.slides._sldIdLst))
