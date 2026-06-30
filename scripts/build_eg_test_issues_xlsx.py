# -*- coding: utf-8 -*-
"""EG品質確認テスト(EG-1/EG-2) 課題整理 Excel生成。

入力 : scratchpad/eg_issues.json  {master:[...], schedule:[...]}
出力 : outputs/EG品質確認テスト_課題整理_20260630.xlsx

シート:
  1. 課題マスタ        全課題（分類・制約・対応方針・期限・担当・RD事前提出・優先度・出典）
  2. RD事前提出依頼    rd_teishutsu=要 を抽出（明日の品質定例で提示する依頼一覧）
  3. スケジュール      テスト工程とSA工程のマイルストーン
  4. 運転課題          大分類=運転 を抽出（上がった運転課題の整理）

体裁ルール（修正ログ正本順守）:
  ・モノクロ（白地・黒文字・薄グレー見出し・細罫線）。色付き装飾・無駄な装飾なし。
  ・マス結合しない。数式を置かない（説明セルも=始まりにしない）。
  ・造語・独自軸・→記号書きをしない。AI痕跡（作成者表記等）を残さない。
"""
import json
import os

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(ROOT, "outputs")
SCRATCH = "/tmp/claude-0/-home-user-Base/78076d83-8708-5756-a6f1-f728259734b3/scratchpad"
os.makedirs(OUT, exist_ok=True)

FONT = "游ゴシック"
HEADER_FILL = "D9D9D9"   # 見出し（薄グレー）
GROUP_FILL = "F2F2F2"    # 区切り/凡例の薄い陰
WHITE = "FFFFFF"
INK = "000000"
BORDER_CLR = "808080"

thin = Side(style="thin", color=BORDER_CLR)
border = Border(left=thin, right=thin, top=thin, bottom=thin)


def fill(c):
    return PatternFill("solid", fgColor=c)


def est_row_height(values, widths):
    """セル内容と列幅からおおまかな行高を見積もる。"""
    lines = 1
    for val, w in zip(values, widths):
        s = "" if val is None else str(val)
        cpl = max(4, int(w * 1.7))  # 1文字≒全角想定の概算
        sub = 0
        for para in s.split("\n"):
            sub += max(1, -(-len(para) // cpl))
        lines = max(lines, sub)
    return min(170, 15 * lines + 4)


def write_sheet(ws, headers, widths, rows, note=None):
    r = 1
    if note:
        c = ws.cell(row=r, column=1, value=note)
        c.font = Font(name=FONT, size=9, color="595959")
        c.alignment = Alignment(horizontal="left", vertical="center", wrap_text=False)
        r += 1
    head_row = r
    for j, h in enumerate(headers, start=1):
        c = ws.cell(row=head_row, column=j, value=h)
        c.font = Font(name=FONT, size=10, bold=True, color=INK)
        c.fill = fill(HEADER_FILL)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = border
    r = head_row + 1
    for row in rows:
        for j, val in enumerate(row, start=1):
            c = ws.cell(row=r, column=j, value=val)
            c.font = Font(name=FONT, size=10, color=INK)
            c.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
            c.border = border
        ws.row_dimensions[r].height = est_row_height(row, widths)
        r += 1
    for j, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = ws.cell(row=head_row + 1, column=1)
    last_col = get_column_letter(len(headers))
    ws.auto_filter.ref = "A%d:%s%d" % (head_row, last_col, max(head_row, r - 1))
    ws.sheet_view.showGridLines = False


def g(d, k, default=""):
    v = d.get(k, default)
    return default if v is None else v


def main():
    with open(os.path.join(SCRATCH, "eg_issues.json"), encoding="utf-8") as f:
        data = json.load(f)
    master = data.get("master", [])
    schedule = data.get("schedule", [])

    wb = Workbook()

    # ---- 1. 課題マスタ ----
    ws1 = wb.active
    ws1.title = "課題マスタ"
    h1 = ["No", "大分類", "小分類", "課題・論点", "制約", "対応方針",
          "期限/マイルストーン", "担当", "RD事前提出", "優先度", "出典・備考"]
    w1 = [4, 8, 14, 30, 26, 30, 16, 12, 9, 7, 22]
    rows1 = []
    for i, m in enumerate(master, start=1):
        bikou = g(m, "shutten")
        if g(m, "bikou"):
            bikou = (bikou + " / " + g(m, "bikou")).strip(" /")
        rows1.append([
            g(m, "no", i), g(m, "bunrui_dai"), g(m, "bunrui_sho"), g(m, "kadai"),
            g(m, "seiyaku"), g(m, "houshin"), g(m, "kigen"), g(m, "tantou"),
            g(m, "rd_teishutsu", "-"), g(m, "yusen"), bikou,
        ])
    write_sheet(ws1, h1, w1, rows1,
                note="EG品質確認テスト(EG-1/EG-2) 課題整理  2026-06-30 C2定例反映")

    # ---- 2. RD事前提出依頼 ----
    ws2 = wb.create_sheet("RD事前提出依頼")
    h2 = ["No", "分類", "RDに事前提出/事前検証してほしい事項", "対応方針/狙い",
          "期限", "担当", "優先度"]
    w2 = [4, 14, 36, 32, 16, 14, 7]
    rows2 = []
    n = 0
    for m in master:
        if str(g(m, "rd_teishutsu", "-")).strip() == "要":
            n += 1
            rows2.append([
                n, g(m, "bunrui_dai") + "/" + g(m, "bunrui_sho"),
                g(m, "kadai"), g(m, "houshin"), g(m, "kigen"),
                g(m, "tantou"), g(m, "yusen"),
            ])
    write_sheet(ws2, h2, w2, rows2,
                note="明日の品質定例で提示：RD(基礎研含む)に事前に出してほしいデータ・検証")

    # ---- 3. スケジュール ----
    ws3 = wb.create_sheet("スケジュール")
    h3 = ["No", "マイルストーン", "時期/期限", "担当", "前提・先行", "備考"]
    w3 = [4, 34, 16, 14, 30, 30]
    rows3 = []
    for i, s in enumerate(schedule, start=1):
        rows3.append([i, g(s, "milestone"), g(s, "target"), g(s, "owner"),
                      g(s, "depends"), g(s, "note")])
    write_sheet(ws3, h3, w3, rows3,
                note="テスト工程とSA工程（課1次SA→部1次SA→課3次SA）のマイルストーン")

    # ---- 4. 運転課題 ----
    ws4 = wb.create_sheet("運転課題")
    h4 = ["No", "小分類", "課題・論点", "制約", "対応方針", "期限", "担当", "優先度"]
    w4 = [4, 22, 30, 26, 30, 16, 14, 7]
    rows4 = []
    n = 0
    for m in master:
        if g(m, "bunrui_dai") == "運転":
            n += 1
            rows4.append([n, g(m, "bunrui_sho"), g(m, "kadai"), g(m, "seiyaku"),
                          g(m, "houshin"), g(m, "kigen"), g(m, "tantou"), g(m, "yusen")])
    write_sheet(ws4, h4, w4, rows4,
                note="上がった運転課題（プロセス・タンク繰り(製販/オフタンク)・生産バランス）")

    out = os.path.join(OUT, "EG品質確認テスト_課題整理_20260630.xlsx")
    wb.save(out)
    print("saved:", out)
    print("master rows:", len(rows1), "| RD依頼:", len(rows2),
          "| schedule:", len(rows3), "| 運転課題:", len(rows4))


if __name__ == "__main__":
    main()
