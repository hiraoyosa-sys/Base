# -*- coding: utf-8 -*-
"""DEG着色シナリオ 横フロー整理（4段化）Excel生成。

横＝プロセスフロー、縦＝4段（残った仮説／残した理由／落とした仮説／落とした理由）。
下部にリン酸の効き場所と一本化した筋を付す。
"""
import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

OUT = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "outputs")
os.makedirs(OUT, exist_ok=True)

DATE = "2026-06-15"

# ---- 配色 ----
C_TITLE   = "1F3864"  # 濃紺
C_HEAD    = "2E5496"  # 工程ヘッダ
C_HEAD_HL = "9E3A26"  # DEG塔(本体)強調
C_KEPT    = "DDEBF7"  # 残った仮説 薄青
C_REASON  = "E2EFDA"  # 残した理由 薄緑
C_DROP    = "F2F2F2"  # 落とした仮説 薄グレー
C_DROPWHY = "D9D9D9"  # 落とした理由 グレー
C_SECT    = "44546A"  # セクション帯
WHITE = "FFFFFF"

thin = Side(style="thin", color="BFBFBF")
border = Border(left=thin, right=thin, top=thin, bottom=thin)

def fill(c):
    return PatternFill("solid", fgColor=c)

# ---- 工程（横） ----
# (見出し, ハイライトか)
STEPS = [
    ("反応系\n(EO反応器)", False),
    ("EG濃縮系\n(パージの場)", False),
    ("脱水塔\n(乾いた条件)", False),
    ("MEG塔\n(C-1501/1502)", False),
    ("DEG塔(C-1503)\n★着色の本体", True),
    ("製品タンク\n(出荷・船・ローリー)", False),
]

# 各工程 4段データ: [残った仮説, 残した理由, 落とした仮説, 落とした理由]
DATA = [
    # 反応系
    [
        "ALD(A/F-ALD)が増加。水が多く遊離のまま。",
        "EO触媒劣化で原料ALD増は事実一致。従来も微量(1ppm以下)常在。",
        "（なし）",
        "—",
    ],
    # EG濃縮系
    [
        "大半はパージで系外へ。一部が隠れ形(アセタール等)で逃げ切る。",
        "A-ALD等はここでパージが通常。原料ALD増で逃げ切る量が増え顕在化、と整合(要検証)。",
        "「全量パージされ下流に来ない」を否定。",
        "実際に下流でALDが見えている。",
    ],
    # 脱水塔
    [
        "アセタール化で重質化し下流まで運ばれる。",
        "水が少なくアセタール側に寄る。重質化すれば届く。",
        "× 軽質ALDのままDEGへ抜ける。",
        "軽質ALDは物性上そのままDEGに来ない(高橋)。運ぶには重質化が必要。",
    ],
    # MEG塔
    [
        "熱・アルカリで隠れ形が分解しALDに戻る(C-1502付近で湧いて見える)。",
        "バランス上C-1502付近でA/F-ALDが生成するように見える(C-1501ボトムは過去同水準=MEG塔単独生成でない)。隠れ形分解で説明可。",
        "× MEG塔そのものが着色の場。",
        "C-1501/1502開放(6/12)→サビ・汚れなし。触媒(活性サビ)の場でない。",
    ],
    # DEG塔
    [
        "高pH＋高温＋活性サビが揃い、アルドール縮合→脱水で共役ポリエナール(300/450nm)。塩基触媒・酸素フリー。",
        "①酸素フリーで進む=酸素で止まる事実(F2)と整合 ②C-1503トップに活性サビ(黒異物・XRF98.5%鉄)実在 ③高pH＋高温が揃う唯一の塔 ④A-ALD自己縮合は共役3-4まで実測(9連結は微量だが微量で色がつく) ⑤活性サビはルイス酸点として働き、高pH(塩基)と組んだ酸塩基二元触媒で脱水(発色)を加速。電子移動・酸素を要さずF2と両立し、高pH＋活性サビが揃う唯一塔=C-1503の根拠を強める。",
        "× 鉄＋酸素のラジカル重合が主役。\n△ ルートB(F+A交差→アクロレイン)は残す(要検証)。",
        "ラジカル重合は酸素が開始剤→酸素で止まる事実と逆、鉄サビ単独で着色せず(RD)。\nルートBは450nm直接証拠なし・寄与度不明で要検証(否定はせず)。\n※落とすのは『ラジカル開始剤としての鉄』。『ルイス酸触媒としての鉄』は残す(残した理由⑤)。",
    ],
    # 製品タンク
    [
        "残った前駆体が保管中に成長→450nm。塩基触媒で酸素なしに進む(縮合/アクロレインのアニオン重合)。",
        "本管DEGは27日でAPHA5以下・無変質、タンク品で450成長=成長はタンク(滞留)で起きる。N2・遮光でも進行=酸素不要。塩基触媒経路は酸素フリー。",
        "× タンクでラジカル重合して着色。",
        "酸素雰囲気(船・ローリー)で悪化が止まる事実と逆(ラジカル重合は酸素が開始剤)。",
    ],
]

ROW_LABELS = ["残った仮説（本流）", "残した理由（なぜ生き残るか）",
              "落とした仮説", "落とした理由（なぜ消えるか＝事実紐付け）"]
ROW_FILLS  = [C_KEPT, C_REASON, C_DROP, C_DROPWHY]

wb = Workbook()
ws = wb.active
ws.title = "フロー整理"
ws.sheet_view.showGridLines = False

NSTEP = len(STEPS)
# 列: A=段ラベル, B..=工程
LBL_W = 22
STEP_W = 30
ws.column_dimensions["A"].width = LBL_W
for i in range(NSTEP):
    ws.column_dimensions[get_column_letter(2 + i)].width = STEP_W

last_col = 1 + NSTEP
last_letter = get_column_letter(last_col)

r = 1
# タイトル
ws.merge_cells(f"A{r}:{last_letter}{r}")
c = ws.cell(r, 1, "DEG着色シナリオ　横フロー整理")
c.font = Font(name="IPAGothic", size=16, bold=True, color=WHITE)
c.fill = fill(C_TITLE)
c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
ws.row_dimensions[r].height = 30
r += 1
# 日付帯
ws.merge_cells(f"A{r}:{last_letter}{r}")
c = ws.cell(r, 1, DATE)
c.font = Font(name="IPAGothic", size=10, color="595959")
c.alignment = Alignment(horizontal="right", vertical="center", indent=1)
ws.row_dimensions[r].height = 16
r += 1

HEADER_ROW = r
# 工程ヘッダ行
lab = ws.cell(r, 1, "工程フロー →")
lab.font = Font(name="IPAGothic", size=10, bold=True, color=WHITE)
lab.fill = fill(C_SECT)
lab.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
lab.border = border
for i, (name, hl) in enumerate(STEPS):
    cell = ws.cell(r, 2 + i, name)
    cell.font = Font(name="IPAGothic", size=10.5, bold=True, color=WHITE)
    cell.fill = fill(C_HEAD_HL if hl else C_HEAD)
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border = border
ws.row_dimensions[r].height = 40
r += 1

# 4段
TIER_HEIGHTS = [72, 128, 58, 118]
for t in range(4):
    lab = ws.cell(r, 1, ROW_LABELS[t])
    lab.font = Font(name="IPAGothic", size=9.5, bold=True, color="333333")
    lab.fill = fill(ROW_FILLS[t])
    lab.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True, indent=1)
    lab.border = border
    for i in range(NSTEP):
        text = DATA[i][t]
        cell = ws.cell(r, 2 + i, text)
        cell.font = Font(name="IPAGothic", size=9, color="222222")
        cell.fill = fill(ROW_FILLS[t])
        cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True, indent=1)
        cell.border = border
    ws.row_dimensions[r].height = TIER_HEIGHTS[t]
    r += 1

r += 1
# ---- リン酸の効き場所 ----
ws.merge_cells(f"A{r}:{last_letter}{r}")
c = ws.cell(r, 1, "リン酸の効き場所（絞り込み済み・確定未）")
c.font = Font(name="IPAGothic", size=11, bold=True, color=WHITE)
c.fill = fill(C_SECT)
c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
ws.row_dimensions[r].height = 22
r += 1
phos = [
    ("効き方", "リン酸鉄被膜が鉄表面サイトを封鎖→鉄の触媒作用を停止。redoxサイクル封じに限定せず、ルイス酸触媒(ルートA加速)としての鉄でも成立しF2と両立。"),
    ("ケースA", "C-1503でP検出 → DEG塔で金属(活性サビ)不活性化＝現行説成立。本線。"),
    ("ケースB", "C-1503でP無し → 効いた場所は上流(MEG/C-1502)で前駆体生成抑制（高橋説）。着色現場=DEG塔とテンション。"),
    ("決め手①", "C-1503ミドル充填部の壁面P分析（足場後採取）。リン入れてMEG-UV効かず=上流説の材料(F8)も突き合わせ。"),
    ("決め手②", "活性黒サビ有/無で酸素フリー着色速度を比較。加速ならルイス酸触媒を直接証明（redox・酸素から分離）。"),
]
for k, v in phos:
    ws.cell(r, 1, k).font = Font(name="IPAGothic", size=9.5, bold=True, color="333333")
    ws.cell(r, 1).fill = fill(C_DROP)
    ws.cell(r, 1).alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.cell(r, 1).border = border
    ws.merge_cells(f"B{r}:{last_letter}{r}")
    cc = ws.cell(r, 2, v)
    cc.font = Font(name="IPAGothic", size=9.5, color="222222")
    cc.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True, indent=1)
    cc.border = border
    for cidx in range(2, last_col + 1):
        ws.cell(r, cidx).border = border
    ws.row_dimensions[r].height = 30
    r += 1

r += 1
# ---- 一本化した筋 ----
ws.merge_cells(f"A{r}:{last_letter}{r}")
c = ws.cell(r, 1, "一本化した筋")
c.font = Font(name="IPAGothic", size=11, bold=True, color=WHITE)
c.fill = fill(C_SECT)
c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
ws.row_dimensions[r].height = 22
r += 1
ws.merge_cells(f"A{r}:{last_letter}{r}")
suji = ("原料ALD増 → 大半パージ・一部隠れ形で逃げる → 脱水塔以降で重質化して運搬 → "
        "MEG塔系で分解しALDに戻り「湧いて見える」 → C-1503(高pH・高温・活性サビ＝酸塩基二元触媒)でアルドール縮合→共役ポリエナール → "
        "タンクで酸素なしに成長 → 450nm着色。\n"
        "動かない軸＝DEGに活性アルデヒドを入れない（IERのガードがなく活性ALDが残ると着色物質に化ける）。")
c = ws.cell(r, 1, suji)
c.font = Font(name="IPAGothic", size=10, color="222222")
c.fill = fill(C_REASON)
c.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True, indent=1)
ws.row_dimensions[r].height = 64
for cidx in range(1, last_col + 1):
    ws.cell(r, cidx).border = border
r += 1

# 印刷設定
ws.print_title_rows = f"{HEADER_ROW}:{HEADER_ROW}"
ws.page_setup.orientation = "landscape"
ws.page_setup.fitToWidth = 1
ws.page_setup.fitToHeight = 0
ws.sheet_properties.pageSetUpPr.fitToPage = True
ws.page_margins.left = ws.page_margins.right = 0.3
ws.page_margins.top = ws.page_margins.bottom = 0.4

# ====== シート2：塔ごとの場所事実 ======
ws2 = wb.create_sheet("塔ごとの場所事実")
ws2.sheet_view.showGridLines = False
ws2.column_dimensions["A"].width = 8
ws2.column_dimensions["B"].width = 90
r = 1
ws2.merge_cells(f"A{r}:B{r}")
c = ws2.cell(r, 1, "確認された事実（F番号）")
c.font = Font(name="IPAGothic", size=14, bold=True, color=WHITE)
c.fill = fill(C_TITLE)
c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
ws2.row_dimensions[r].height = 28
r += 1
FACTS = [
    ("F1", "採取直後無色、保管中(タンク)に着色進行。"),
    ("F2", "酸素雰囲気(船・ローリー)で悪化が止まる。N2・遮光でも進行。【最重要絞り込み】"),
    ("F3", "300＋450nm=長い共役系(ポリエナール)。"),
    ("F4", "C-1503トップに黒異物・サビ(XRF98.5%鉄)。C-1501/1502はサビ・汚れなし(6/12)。"),
    ("F5", "鉄サビ単独では着色せず(RD保管・蒸留試験)。"),
    ("F6", "A-ALD縮合は共役3-4まで実測。9連結は微量だが微量で色がつく。"),
    ("F7", "本管DEG27日でAPHA5以下・無変質、タンク品で450成長。"),
    ("F8", "着色に効いたのはリン酸のみ。ただしリン入れてもMEG-UV効かず。"),
    ("F9", "バランス上C-1502付近でALD生成に見える。C-1501ボトムは過去同水準。"),
    ("F10", "軽質ALDは物性上そのままDEGに来ない(高橋)。運ぶには重質化必要。"),
]
for k, v in FACTS:
    kc = ws2.cell(r, 1, k)
    kc.font = Font(name="IPAGothic", size=10, bold=True, color="FFFFFF")
    kc.fill = fill(C_HEAD)
    kc.alignment = Alignment(horizontal="center", vertical="center")
    kc.border = border
    vc = ws2.cell(r, 2, v)
    hl = (k == "F2")
    vc.font = Font(name="IPAGothic", size=10, color="222222", bold=hl)
    vc.fill = fill(C_REASON if hl else WHITE)
    vc.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True, indent=1)
    vc.border = border
    ws2.row_dimensions[r].height = 30
    r += 1

r += 1
ws2.merge_cells(f"A{r}:B{r}")
c = ws2.cell(r, 1, "塔ごとの場所事実（追記枠）")
c.font = Font(name="IPAGothic", size=12, bold=True, color=WHITE)
c.fill = fill(C_SECT)
c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
ws2.row_dimensions[r].height = 24
r += 1
PLACE = [
    ("C-1501/1502", "開放(6/12)でサビ・汚れなし。活性サビの場でない＝着色の場ではない。"),
    ("C-1503トップ", "黒異物・活性サビ(XRF98.5%鉄)実在。高pH＋高温が揃う唯一の塔＝着色の本体。"),
    ("C-1503ミドル充填部", "壁面P分析を足場後に採取予定（リン酸の効き場所の決め手）。"),
    ("製品タンク", "滞留で450nm成長。N2・遮光でも進行＝酸素不要（塩基触媒経路）。"),
]
for k, v in PLACE:
    kc = ws2.cell(r, 1, k)
    kc.font = Font(name="IPAGothic", size=9.5, bold=True, color="333333")
    kc.fill = fill(C_DROP)
    kc.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    kc.border = border
    ws2.column_dimensions["A"].width = 16
    vc = ws2.cell(r, 2, v)
    vc.font = Font(name="IPAGothic", size=9.5, color="222222")
    vc.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True, indent=1)
    vc.border = border
    ws2.row_dimensions[r].height = 30
    r += 1

ws2.page_setup.orientation = "landscape"
ws2.page_setup.fitToWidth = 1
ws2.page_setup.fitToHeight = 0
ws2.sheet_properties.pageSetUpPr.fitToPage = True

# ---- メタデータ（作成者）----
wb.properties.creator = "hirao kazuaki/0465811/平尾　一陽"
wb.properties.lastModifiedBy = "hirao kazuaki/0465811/平尾　一陽"
wb.properties.title = "DEG着色シナリオ 横フロー整理"
wb.properties.description = ""
wb.properties.keywords = ""

path = os.path.join(OUT, "DEG着色_フロー整理.xlsx")
wb.save(path)
print("saved:", path)
