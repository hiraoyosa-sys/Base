#!/usr/bin/env python3
"""HAZchart簡易評価ワークブックを生成する。

出力: outputs/HAZchart計算ベース_rev20260721.xlsx
シート: 読み方 / 計算原理 / 確率DB / シナリオ計算 / 例題_V1204 / リスクマトリックス

確率DBの数値は、アップロードされた project.pha（PHA_Organizer 3.1.1）の
テンプレート設定値を scripts/extract_pha_values.py で抽出したもの（2026-07-21）。
体裁: セル結合なし・標準色のみ・入力セルは黄色塗り＋青字・数式に @ を使わない。
"""
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

FONT = 'Yu Gothic'
F_BASE = Font(name=FONT, size=10)
F_BOLD = Font(name=FONT, size=10, bold=True)
F_TITLE = Font(name=FONT, size=14, bold=True)
F_SECTION = Font(name=FONT, size=11, bold=True)
F_INPUT = Font(name=FONT, size=10, color='0000FF')          # 入力値=青字
F_NOTE = Font(name=FONT, size=9, color='595959')
F_WARN = Font(name=FONT, size=10, bold=True, color='FF0000')
FILL_HDR = PatternFill('solid', fgColor='D9D9D9')
FILL_IN = PatternFill('solid', fgColor='FFFF00')            # 触るセル=黄
FILL_SEC = PatternFill('solid', fgColor='F2F2F2')
FILL_RED = PatternFill('solid', fgColor='FF6666')
FILL_YEL = PatternFill('solid', fgColor='FFFF99')
FILL_GRN = PatternFill('solid', fgColor='99CC99')
THIN = Side(style='thin', color='808080')
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
SCI = '0.00E+00'

# ---- project.pha 抽出値（一次データ / 2026-07-21 抽出）----------------------
SRC = 'project.pha（PHA_Organizer 3.1.1）設定値を抽出 2026-07-21'
DB = [
    # 分類, 要素名, 故障モード, λ(/hr), T(h), デマンド確率(/d), 備考
    ('コモン事象', 'DCS', '故障', 5.28e-10, None, 1.94e-07,
     'DCSは各制御ループ・検出システムに共通で入る（コモン事象）'),
    ('検出端', '温度計', '故障', 9.7e-05, 720, 0.035, '月1回点検（T=720h）の待機系換算'),
    ('検出端', '流量計', '故障', 1.2e-04, 720, 0.0433, '月1回点検（T=720h）の待機系換算'),
    ('検出端', '液面計', '故障', 1.9e-04, 720, 0.0685, '月1回点検（T=720h）の待機系換算'),
    ('検出端', '圧力計', '故障', 2.2e-04, 720, 0.0793, '月1回点検（T=720h）の待機系換算'),
    ('検出端', '分析計', '故障', 9.7e-04, None, 0.093, ''),
    ('警報', 'アラーム', '故障', 7.7e-07, None, 9.3e-06, ''),
    ('ロジック', 'リレー回路', '故障', 2.0e-06, None, 6.9e-04, 'インターロックの論理回路'),
    ('操作端', '空気調整弁', '故障', 3.6e-06, None, 0.0022, '制御弁・遮断弁の駆動部'),
    ('操作端', '電磁弁', '故障', 4.9e-05, None, 0.0029, ''),
    ('機械', 'ポンプ', '故障', 4.4e-05, None, 0.002, ''),
    ('機械', 'コンプレッサー', '故障', 2.8e-05, None, 0.0, 'デマンド確率は未設定（0）'),
    ('機械', 'チェッキ弁', '故障', 3.2e-06, None, 0.0022, ''),
    ('機械', '安全弁', '故障', 1.7e-06, None, 2.2e-04, ''),
    ('ヒューマンエラー', 'ヒューマンエラー', 'アラームに気づかない', None, None, 1.0e-04, ''),
    ('ヒューマンエラー', 'ヒューマンエラー', '発見したが対処できない（時間余裕10分）', None, None, 0.3, ''),
    ('ヒューマンエラー', 'ヒューマンエラー', '発見したが対処できない（時間余裕1時間）', None, None, 0.006, ''),
    ('ヒューマンエラー', 'ヒューマンエラー', '発見したが対処できない（時間余裕1日）', None, None, 3.0e-04, ''),
]
DB_TOP = 4                       # データ開始行
DB_BOTTOM = DB_TOP + len(DB) + 5  # 追記用の空行を5行確保
KEY_RNG = f"'確率DB'!$B${DB_TOP}:$B${DB_BOTTOM}"
LAM_RNG = f"'確率DB'!$E${DB_TOP}:$E${DB_BOTTOM}"
PFD_RNG = f"'確率DB'!$G${DB_TOP}:$G${DB_BOTTOM}"


def key(comp, mode):
    return f'{comp}／{mode}' if mode != '故障' else f'{comp}／故障'


def style(ws, addr, value=None, font=F_BASE, fill=None, fmt=None, wrap=False,
          border=False, align=None):
    c = ws[addr]
    if value is not None:
        c.value = value
    c.font = font
    if fill:
        c.fill = fill
    if fmt:
        c.number_format = fmt
    if wrap or align:
        c.alignment = Alignment(wrap_text=wrap, horizontal=align, vertical='center')
    if border:
        c.border = BORDER
    return c


def widths(ws, spec):
    for col, w in spec.items():
        ws.column_dimensions[col].width = w


def lookup_pfd(cell):
    return f'=IFERROR(INDEX({PFD_RNG},MATCH({cell},{KEY_RNG},0)),0)'


def lookup_lam(cell):
    return f'=IFERROR(INDEX({LAM_RNG},MATCH({cell},{KEY_RNG},0)),0)'


wb = Workbook()

# ============================================================ 1. 読み方
ws = wb.active
ws.title = '読み方'
widths(ws, {'A': 3, 'B': 26, 'C': 95})
style(ws, 'B2', 'HAZchart 簡易評価ワークブック（計算原理＋確率DB＋シナリオシミュレーション）', F_TITLE)
style(ws, 'B3', '作成 2026-07-21。確率DBの出典: ' + SRC, F_NOTE)
rows = [
    ('目的', 'PHA_Organizerを使わずに、HAZchartの確率計算の仕組みを理解し、'
             'シナリオを組んだときの事故頻度と安全対策（防護層の追加・削除）の効果をExcel上で素早く定量評価する。'),
    ('位置づけ', '簡易評価用（オーダー確認・感度検討）。正式なリスク評価はHAZchart解析基準（C-2-3）に'
                 '基づきPHA_Organizerで実施する。本ブックの結果はその代替にはならない。'),
    ('シート構成', '「計算原理」…計算の仕組みの説明。 「確率DB」…既存テンプレートの故障率・デマンド確率の一覧（別シート切り分け）。'
                   ' 「シナリオ計算」…シナリオを組んで頻度を計算するシミュレーター。'
                   ' 「例題_V1204」…アップロードされたプロジェクトの再現例と検算。 「リスクマトリックス」…頻度・影響度の判定表（仮置き）。'),
    ('使い方(1)', '「シナリオ計算」で起因事象の構成要素をドロップダウンから選ぶ。制御弁の開度異常は'
                  '「検出器＋DCS＋操作端」の3要素を選ぶと制御ループ故障の合計になる。'),
    ('使い方(2)', '防護層を上から通過する順に並べる。各層の構成要素を選ぶと、その層が失敗する確率（PFD）が自動計算される。'),
    ('使い方(3)', '「有効」列を1から0に切り替えると、その層が無い場合の頻度が即座に分かる。'
                  '逆に空いている行に層を追加すれば、インターロック等を1つ追加した効果が定量で見える。'),
    ('凡例', '黄色塗り＝触って試すセル（青字）。 白地の黒字＝自動計算（触らない）。 灰色＝見出し。'),
]
r = 5
for label, text in rows:
    style(ws, f'B{r}', label, F_BOLD, fill=FILL_SEC, border=True, align='left')
    style(ws, f'C{r}', text, F_BASE, wrap=True, border=True)
    ws.row_dimensions[r].height = 30
    r += 1
style(ws, f'B{r + 1}', '注意（重要）', F_WARN)
notes = [
    '同じ要素（例: DCS、同一の検出器）が複数の層に入る場合、単純な掛け算は頻度を過小評価する（コモン事象）。'
    '同一要素は1つの層にだけ入れること。厳密な扱いはPHA_Organizerのミニマルカットセット計算で行う。',
    '層のPFDは構成要素の和で近似している（OR近似）。合計が0.1を超えるあたりから過大側にずれる（安全側）。',
    'リスクマトリックスの区分・ゾーン割付は一般形の仮置き。C-2-3の正式な表と照合して確定すること（要確認）。',
]
r += 2
for t in notes:
    style(ws, f'C{r}', '・' + t, F_BASE, wrap=True)
    ws.row_dimensions[r].height = 28
    r += 1

# ============================================================ 2. 計算原理
ws = wb.create_sheet('計算原理')
widths(ws, {'A': 3, 'B': 34, 'C': 22, 'D': 95})
style(ws, 'B2', 'HAZchartの確率計算の原理', F_TITLE)

r = 4
style(ws, f'B{r}', '1. 解析の4ステップ', F_SECTION, fill=FILL_SEC)
steps = [
    ('ステップ1', 'What-ifで事故シナリオを作る（起因事象＝機器の故障モード。防護層は無いものとして展開）'),
    ('ステップ2', 'HAZchartを描く（状態変化の認知手段、回復動作、失敗したときの次の状態、を繰り返す）'),
    ('ステップ3', 'FT（フォールトツリー）に変換する（最終事象を頂上に置き、状態変化と防護失敗をANDで結ぶ）'),
    ('ステップ4', '故障率・ヒューマンエラー確率を入れて最終事象の発生頻度を計算し、リスクマトリックスで評価する'),
]
for label, text in steps:
    r += 1
    style(ws, f'B{r}', label, F_BOLD, border=True, fill=FILL_SEC)
    style(ws, f'D{r}', text, F_BASE, border=True, wrap=True)

r += 2
style(ws, f'B{r}', '2. 確率は2種類ある', F_SECTION, fill=FILL_SEC)
for label, text in [
    ('時間確率 λ（/hr）', '運転中いつでも起こりうる故障の発生率。起因事象に使う。年間頻度＝λ×年間稼働時間H。'),
    ('デマンド確率（/d）', '「要求されたときに動かない確率」。検出・警報・人の対処・インターロック・安全弁など防護層に使う。'),
]:
    r += 1
    style(ws, f'B{r}', label, F_BOLD, border=True, fill=FILL_SEC)
    style(ws, f'D{r}', text, F_BASE, border=True, wrap=True)

r += 2
style(ws, f'B{r}', '3. 基本式（この4つだけ）', F_SECTION, fill=FILL_SEC)
formulas = [
    ('起因事象の頻度', 'F0（/年） ＝ λ合計（/hr） × 年間稼働時間H（h/年）。'
     '制御ループの開度異常などは、検出器＋DCS＋操作端のλの和（ORゲート）。'),
    ('最終事象の頻度', 'F（/年） ＝ F0 × PFD層1 × PFD層2 × …。'
     '防護層は「すべて失敗したとき」だけ最終事象に至るのでANDゲート＝掛け算。'),
    ('層のPFD（失敗確率）', '層PFD ≒ 構成要素のデマンド確率の和（ORゲートの近似）。'
     '厳密には 1−(1−P1)×(1−P2)×… だが、Pが小さければ和でよい（和は少し大きめ＝安全側）。'),
    ('待機系の換算', '点検周期Tで点検される待機機器は デマンド確率 ＝ λ × T ÷ 2。'
     '（故障してから次の点検までの平均放置時間がT/2のため。月1回点検ならT=720h）'),
]
for label, text in formulas:
    r += 1
    style(ws, f'B{r}', label, F_BOLD, border=True, fill=FILL_SEC)
    style(ws, f'D{r}', text, F_BASE, border=True, wrap=True)
    ws.row_dimensions[r].height = 28

r += 2
style(ws, f'B{r}', '4. 層の中身の組み立て方（FT展開の型）', F_SECTION, fill=FILL_SEC)
for label, text in [
    ('人が対処する層', '層PFD ＝ 発見手段の故障（検出器＋DCS＋アラーム） ＋ アラームに気づかない ＋ 発見したが対処できない。'
     '「対処できない」は時間余裕で値が変わる（10分: 0.3 / 1時間: 0.006 / 1日: 0.0003）。'),
    ('自動の層（インターロック）', '層PFD ＝ 検出器 ＋ ロジック（リレー回路） ＋ 操作端（電磁弁・空気調整弁）。人は介在しない。'),
    ('機械式の層', '安全弁・チェッキ弁などは単体のデマンド確率をそのまま使う。'),
]:
    r += 1
    style(ws, f'B{r}', label, F_BOLD, border=True, fill=FILL_SEC)
    style(ws, f'D{r}', text, F_BASE, border=True, wrap=True)
    ws.row_dimensions[r].height = 28

r += 2
style(ws, f'B{r}', '5. 生きた検算（式が入っています）', F_SECTION, fill=FILL_SEC)
r += 1
style(ws, f'B{r}', '待機系換算の確認: 温度計 λ×T÷2', F_BASE, border=True)
style(ws, f'C{r}', f"=INDEX({LAM_RNG},MATCH(\"温度計／故障\",{KEY_RNG},0))*720/2",
      F_BASE, border=True, fmt=SCI)
style(ws, f'D{r}', '確率DBの登録値0.035と一致する（登録値がこの式で作られていることの確認）', F_NOTE, wrap=True)
r += 1
style(ws, f'B{r}', 'OR近似の誤差例: P1=0.3, P2=0.08 の和', F_BASE, border=True)
style(ws, f'C{r}', '=0.3+0.08', F_BASE, border=True, fmt='0.000')
style(ws, f'D{r}', '厳密値は 1−(0.7×0.92)＝0.356。和(0.38)は少し大きめに出る＝安全側。', F_NOTE, wrap=True)
r += 1
style(ws, f'B{r}', '起因頻度の例: λ=1.0E-4/hr, H=8760h', F_BASE, border=True)
style(ws, f'C{r}', '=0.0001*8760', F_BASE, border=True, fmt='0.000')
style(ws, f'D{r}', '1年に0.88回起こる、と読む。', F_NOTE)

r += 2
style(ws, f'B{r}', '6. コモン事象（掛け算の落とし穴）', F_SECTION, fill=FILL_SEC)
r += 1
style(ws, f'D{r}', '同じ基本事象（例: DCS故障）が複数の層に現れると、掛け算では同じ故障を2回掛けてしまい頻度を過小評価する。'
      'PHA_Organizerはブール代数（ミニマルカットセット）で重複を正しく処理する。'
      '本ブックの簡易計算では、同一要素は1つの層にだけ入れて近似すること。', F_BASE, wrap=True)
ws.row_dimensions[r].height = 40
r += 2
style(ws, f'D{r}', '出典: 計算体系はHAZchart解析基準（C-2-3）とPHA_Organizerの処理に基づく。'
      '数値の関係（λ×T/2＝登録デマンド確率）は' + SRC + 'で確認済み。', F_NOTE, wrap=True)

# ============================================================ 3. 確率DB
ws = wb.create_sheet('確率DB')
widths(ws, {'A': 12, 'B': 40, 'C': 14, 'D': 34, 'E': 12, 'F': 11, 'G': 14, 'H': 13, 'I': 40, 'J': 46})
style(ws, 'A1', '確率DB（既存テンプレートの設定値）', F_TITLE)
style(ws, 'A2', '出典: ' + SRC + '。追記する場合は下の空行に同じ形式で入力（キー列は「要素名／故障モード」）。', F_NOTE)
hdrs = ['分類', 'キー（選択用）', '要素名', '故障モード', '時間確率λ(/hr)', '点検周期T(h)',
        'デマンド確率(/d)', '検算 λ×T÷2', '出典', '備考']
for i, h in enumerate(hdrs, start=1):
    style(ws, f'{get_column_letter(i)}3', h, F_BOLD, fill=FILL_HDR, border=True, wrap=True, align='center')
r = DB_TOP
for cat, comp, mode, lam, T, pfd, note in DB:
    style(ws, f'A{r}', cat, F_BASE, border=True)
    style(ws, f'B{r}', key(comp, mode), F_BASE, border=True)
    style(ws, f'C{r}', comp, F_BASE, border=True)
    style(ws, f'D{r}', mode, F_BASE, border=True)
    style(ws, f'E{r}', lam if lam is not None else '', F_BASE, border=True, fmt=SCI)
    style(ws, f'F{r}', T if T is not None else '', F_BASE, border=True, fmt='0')
    style(ws, f'G{r}', pfd, F_BASE, border=True, fmt=SCI)
    style(ws, f'H{r}', f'=IF(F{r}="","",E{r}*F{r}/2)', F_BASE, border=True, fmt=SCI)
    style(ws, f'I{r}', SRC, F_NOTE, border=True, wrap=True)
    style(ws, f'J{r}', note, F_NOTE, border=True, wrap=True)
    r += 1
for rr in range(r, DB_BOTTOM + 1):
    for col in 'ABCDEFGHIJ':
        style(ws, f'{col}{rr}', '', F_INPUT, fill=FILL_IN if col in 'ABCDEFG' else None, border=True)
style(ws, f'A{DB_BOTTOM + 1}', '（黄色の行は追記用。λとTを入れれば検算列でデマンド確率の目安を作れる）', F_NOTE)
ws.freeze_panes = 'A4'
wb.defined_names.add(DefinedName('DB_KEYS', attr_text=KEY_RNG))

# ============================================================ 4. シナリオ計算
ws = wb.create_sheet('シナリオ計算')
widths(ws, {'A': 5, 'B': 34, 'C': 24, 'D': 24, 'E': 24, 'F': 24, 'G': 24, 'H': 24,
            'I': 11, 'J': 11, 'K': 11, 'L': 11, 'M': 11, 'N': 11,
            'O': 12, 'P': 12, 'Q': 7, 'R': 12, 'S': 15})
style(ws, 'A1', 'シナリオ計算（シミュレーター）', F_TITLE)
style(ws, 'A2', '黄色セルだけ触る。初期値には例題（FC12開度過少、V-1204レベル上昇）を入れてある。', F_NOTE)

# STEP0
style(ws, 'B4', 'STEP0 前提', F_SECTION, fill=FILL_SEC)
style(ws, 'B5', '年間稼働時間 H（h/年）', F_BASE, border=True)
style(ws, 'C5', 8760, F_INPUT, fill=FILL_IN, border=True, fmt='0')
style(ws, 'D5', '通年連続運転なら8760。定修等を除くなら実稼働時間に変更。', F_NOTE, wrap=True)

# STEP1
style(ws, 'B7', 'STEP1 起因事象（機器の故障）', F_SECTION, fill=FILL_SEC)
style(ws, 'D7', '制御弁の開度異常は「検出器＋DCS＋操作端」の3要素を選ぶ（制御ループ故障の合計になる）', F_NOTE)
style(ws, 'B8', '構成要素（ドロップダウンで選択）', F_BOLD, fill=FILL_HDR, border=True)
style(ws, 'C8', 'λ(/hr)', F_BOLD, fill=FILL_HDR, border=True, align='center')
style(ws, 'D8', 'メモ', F_BOLD, fill=FILL_HDR, border=True)
init_events = [('流量計／故障', 'FC12制御ループの検出器'),
               ('DCS／故障', '同ループのDCS'),
               ('空気調整弁／故障', '同ループの操作端'),
               ('', '')]
for i, (k, memo) in enumerate(init_events):
    rr = 9 + i
    style(ws, f'B{rr}', k, F_INPUT, fill=FILL_IN, border=True)
    style(ws, f'C{rr}', lookup_lam(f'B{rr}'), F_BASE, border=True, fmt=SCI)
    style(ws, f'D{rr}', memo, F_INPUT, fill=FILL_IN, border=True)
style(ws, 'B13', 'λ合計（/hr）', F_BOLD, border=True)
style(ws, 'C13', '=SUM(C9:C12)', F_BOLD, border=True, fmt=SCI)
style(ws, 'B14', '起因事象頻度 F0（/年）', F_BOLD, border=True)
style(ws, 'C14', '=C13*$C$5', F_BOLD, border=True, fmt=SCI)

# STEP2
style(ws, 'B16', 'STEP2 安全防護層（シナリオで通過する順に上から）', F_SECTION, fill=FILL_SEC)
style(ws, 'B17', '各層のPFD＝構成要素のデマンド確率の和（OR近似）。手入力PFDに値を入れるとそちらを優先。'
      '「有効」を0にするとその層なし（未設置）として計算。', F_NOTE)
hdr2 = ['層', '層の名称', '構成要素1', '構成要素2', '構成要素3', '構成要素4', '構成要素5', '構成要素6',
        'PFD1', 'PFD2', 'PFD3', 'PFD4', 'PFD5', 'PFD6', '層PFD(自動)', '手入力PFD', '有効',
        '採用PFD', 'この層まで全て失敗した頻度(/年)']
for i, h in enumerate(hdr2, start=1):
    style(ws, f'{get_column_letter(i)}18', h, F_BOLD, fill=FILL_HDR, border=True, wrap=True, align='center')
ws.row_dimensions[18].height = 30
L_TOP, L_N = 19, 8
init_layers = {
    0: ('LK液面検出で運転員が対処（時間余裕10分）',
        ['液面計／故障', 'DCS／故障', 'アラーム／故障',
         'ヒューマンエラー／アラームに気づかない', 'ヒューマンエラー／発見したが対処できない（時間余裕10分）', '']),
    1: ('OSS（圧力検出で自動遮断）',
        ['圧力計／故障', 'リレー回路／故障', '空気調整弁／故障', '', '', '']),
}
for i in range(L_N):
    rr = L_TOP + i
    name, comps = init_layers.get(i, ('', ['', '', '', '', '', '']))
    style(ws, f'A{rr}', i + 1, F_BASE, border=True, align='center')
    style(ws, f'B{rr}', name, F_INPUT, fill=FILL_IN, border=True, wrap=True)
    for j, comp in enumerate(comps):
        col = get_column_letter(3 + j)          # C..H
        style(ws, f'{col}{rr}', comp, F_INPUT, fill=FILL_IN, border=True, wrap=True)
        pcol = get_column_letter(9 + j)         # I..N
        style(ws, f'{pcol}{rr}', lookup_pfd(f'{col}{rr}'), F_BASE, border=True, fmt=SCI)
    style(ws, f'O{rr}', f'=MIN(1,SUM(I{rr}:N{rr}))', F_BASE, border=True, fmt=SCI)
    style(ws, f'P{rr}', '', F_INPUT, fill=FILL_IN, border=True, fmt=SCI)
    style(ws, f'Q{rr}', 1, F_INPUT, fill=FILL_IN, border=True, align='center')
    style(ws, f'R{rr}', f'=IF($B{rr}="",1,IF($Q{rr}=0,1,IF($P{rr}<>"",$P{rr},$O{rr})))',
          F_BASE, border=True, fmt=SCI)
    prev = '$C$14' if i == 0 else f'S{rr - 1}'
    style(ws, f'S{rr}', f'={prev}*R{rr}', F_BASE, border=True, fmt=SCI)
    ws.row_dimensions[rr].height = 30

# STEP3
R0 = L_TOP + L_N + 1
style(ws, f'B{R0}', 'STEP3 結果', F_SECTION, fill=FILL_SEC)
style(ws, f'B{R0+1}', '最終事象の頻度（/年）', F_BOLD, border=True)
style(ws, f'C{R0+1}', f'=S{L_TOP + L_N - 1}', F_BOLD, border=True, fmt=SCI)
style(ws, f'B{R0+2}', '（/hr換算）', F_BASE, border=True)
style(ws, f'C{R0+2}', f'=C{R0+1}/$C$5', F_BASE, border=True, fmt=SCI)
style(ws, f'B{R0+3}', '頻度レベル（参考・要確認）', F_BASE, border=True)
style(ws, f'C{R0+3}',
      f'=IF(C{R0+1}>=0.1,"E",IF(C{R0+1}>=0.01,"D",IF(C{R0+1}>=0.0001,"C",IF(C{R0+1}>=0.000001,"B","A"))))',
      F_BOLD, border=True, align='center')
style(ws, f'B{R0+4}', '影響度レベル（1〜4を入力）', F_BASE, border=True)
style(ws, f'C{R0+4}', 2, F_INPUT, fill=FILL_IN, border=True, align='center')
style(ws, f'B{R0+5}', 'リスクゾーン（参考・要確認）', F_BASE, border=True)
style(ws, f'C{R0+5}',
      f"=INDEX('リスクマトリックス'!$H$18:$K$22,MATCH(C{R0+3},'リスクマトリックス'!$G$18:$G$22,0),C{R0+4})",
      F_BOLD, border=True, align='center')
style(ws, f'D{R0+3}', '頻度区分・ゾーンはリスクマトリックスシートの仮置き表による。C-2-3の正式表と照合のこと。', F_NOTE, wrap=True)

# What-if
W0 = R0 + 7
style(ws, f'B{W0}', '対策効果の比較（What-if）', F_SECTION, fill=FILL_SEC)
style(ws, f'B{W0+1}', 'ベースケースの頻度（数値を書き留める）', F_BASE, border=True)
style(ws, f'C{W0+1}', '', F_INPUT, fill=FILL_IN, border=True, fmt=SCI)
style(ws, f'D{W0+1}', '現状ケースの頻度をここに数値でメモしてから層を変更すると、下で低減倍率が出る。', F_NOTE, wrap=True)
style(ws, f'B{W0+2}', '低減倍率（ベース÷現ケース）', F_BASE, border=True)
style(ws, f'C{W0+2}', f'=IF(C{W0+1}="","",C{W0+1}/C{R0+1})', F_BOLD, border=True, fmt='0.0')

dv = DataValidation(type='list', formula1='=DB_KEYS', allow_blank=True, showErrorMessage=False)
ws.add_data_validation(dv)
dv.add('B9:B12')
dv.add(f'C{L_TOP}:H{L_TOP + L_N - 1}')
ws.freeze_panes = 'A4'

# ============================================================ 5. 例題_V1204
ws = wb.create_sheet('例題_V1204')
widths(ws, {'A': 3, 'B': 8, 'C': 46, 'D': 40, 'E': 14, 'F': 14, 'G': 60})
style(ws, 'B1', '例題: FC12開度過少によるV-1204レベル上昇（アップロードされたproject.phaの再現）', F_TITLE)
style(ws, 'B2', 'シナリオ（What-if）の流れ', F_SECTION, fill=FILL_SEC)
seq = [
    '1. 起因事象: FC12 開度過少（制御ループ故障）',
    '2. フォーミングが進行してV-1204のレベルが上がる',
    '3. LK液面検出システムで検知し、運転員が回復動作（時間余裕10分）',
    '4. 回復動作に失敗すると、さらにレベルが上がる',
    '5. OSSが作動して自動停止する',
    '6. OSSが不作動なら、最終事象「EO反応器まで水を巻き上げて触媒失活」に至る',
]
r = 3
for s in seq:
    style(ws, f'C{r}', s, F_BASE)
    r += 1

r += 1
style(ws, f'B{r}', '計算の再現（式が入っています。確率DBを参照）', F_SECTION, fill=FILL_SEC)
r += 1
hdr = ['項目', '構成要素', '値', '単位', '備考']
for i, h in enumerate(hdr):
    style(ws, f'{get_column_letter(3 + i)}{r}', h, F_BOLD, fill=FILL_HDR, border=True, align='center')
rows_calc = [
    ('起因事象', '流量計／故障', lookup_lam, '/hr', 'FC12ループの検出器'),
    ('起因事象', 'DCS／故障', lookup_lam, '/hr', ''),
    ('起因事象', '空気調整弁／故障', lookup_lam, '/hr', 'FC12ループの操作端'),
    ('防護層1', '液面計／故障', lookup_pfd, '/d', 'LK液面検出'),
    ('防護層1', 'DCS／故障', lookup_pfd, '/d', ''),
    ('防護層1', 'アラーム／故障', lookup_pfd, '/d', ''),
    ('防護層1', 'ヒューマンエラー／アラームに気づかない', lookup_pfd, '/d', ''),
    ('防護層1', 'ヒューマンエラー／発見したが対処できない（時間余裕10分）', lookup_pfd, '/d', ''),
    ('防護層2', '圧力計／故障', lookup_pfd, '/d', 'OSSの検出（PDK）'),
    ('防護層2', 'リレー回路／故障', lookup_pfd, '/d', 'OSSのロジック'),
    ('防護層2', '空気調整弁／故障', lookup_pfd, '/d', 'OSSの操作端'),
]
first = r + 1
for i, (grp, comp, fn, unit, note) in enumerate(rows_calc):
    rr = first + i
    style(ws, f'C{rr}', grp, F_BASE, border=True)
    style(ws, f'D{rr}', comp, F_BASE, border=True, wrap=True)
    style(ws, f'E{rr}', fn(f'D{rr}'), F_BASE, border=True, fmt=SCI)
    style(ws, f'F{rr}', unit, F_BASE, border=True, align='center')
    style(ws, f'G{rr}', note, F_NOTE, border=True)
lam_rows = f'E{first}:E{first + 2}'
l1_rows = f'E{first + 3}:E{first + 7}'
l2_rows = f'E{first + 8}:E{first + 10}'
r = first + len(rows_calc) + 1
style(ws, f'C{r}', '起因事象頻度 F0 ＝ λ合計 × 8760', F_BOLD, border=True)
style(ws, f'E{r}', f'=SUM({lam_rows})*8760', F_BOLD, border=True, fmt=SCI)
style(ws, f'F{r}', '/年', F_BASE, border=True, align='center')
F0 = f'E{r}'
r += 1
style(ws, f'C{r}', '防護層1のPFD（要素の和）', F_BOLD, border=True)
style(ws, f'E{r}', f'=MIN(1,SUM({l1_rows}))', F_BOLD, border=True, fmt=SCI)
style(ws, f'F{r}', '/d', F_BASE, border=True, align='center')
P1 = f'E{r}'
r += 1
style(ws, f'C{r}', '防護層2のPFD（要素の和）', F_BOLD, border=True)
style(ws, f'E{r}', f'=MIN(1,SUM({l2_rows}))', F_BOLD, border=True, fmt=SCI)
style(ws, f'F{r}', '/d', F_BASE, border=True, align='center')
P2 = f'E{r}'
r += 1
style(ws, f'C{r}', '最終事象頻度 F ＝ F0 × 層1PFD × 層2PFD', F_BOLD, border=True)
style(ws, f'E{r}', f'={F0}*{P1}*{P2}', F_BOLD, border=True, fmt=SCI)
style(ws, f'F{r}', '/年', F_BASE, border=True, align='center')
FE = f'E{r}'
r += 1
style(ws, f'C{r}', '（/hr換算）', F_BASE, border=True)
style(ws, f'E{r}', f'={FE}/8760', F_BASE, border=True, fmt=SCI)
style(ws, f'F{r}', '/hr', F_BASE, border=True, align='center')
FH = f'E{r}'

r += 2
style(ws, f'B{r}', 'PHA_Organizerの計算値との比較', F_SECTION, fill=FILL_SEC)
r += 1
style(ws, f'C{r}', 'PHA_Organizer保存値（project.pha内 FT-0002）', F_BASE, border=True)
style(ws, f'E{r}', 5.24797e-06, F_BASE, border=True, fmt=SCI)
style(ws, f'F{r}', '/hr', F_BASE, border=True, align='center')
style(ws, f'G{r}', '出典: ' + SRC, F_NOTE, border=True)
PHA = f'E{r}'
r += 1
style(ws, f'C{r}', '本シートの近似値 ÷ PHA_Organizer値', F_BASE, border=True)
style(ws, f'E{r}', f'={FH}/{PHA}', F_BASE, border=True, fmt='0.00')
style(ws, f'G{r}', '同オーダー（1未満＝やや小さめ）。差はFTの分岐（PDK・TK経由の経路）と'
      'カットセット処理を簡略化しているため（推測）。オーダー確認用として使う。', F_NOTE, border=True, wrap=True)
ws.row_dimensions[r].height = 30

r += 2
style(ws, f'B{r}', '対策効果の見える化（この例でのWhat-if）', F_SECTION, fill=FILL_SEC)
r += 1
for i, h in enumerate(['ケース', '内容', '頻度(/年)', '現状との比']):
    style(ws, f'{get_column_letter(3 + i)}{r}', h, F_BOLD, fill=FILL_HDR, border=True, align='center')
r += 1
base_row = r
style(ws, f'C{r}', '現状', F_BASE, border=True)
style(ws, f'D{r}', '層1（LK＋運転員）＋層2（OSS）', F_BASE, border=True)
style(ws, f'E{r}', f'={F0}*{P1}*{P2}', F_BASE, border=True, fmt=SCI)
style(ws, f'F{r}', 1, F_BASE, border=True, fmt='0.0')
r += 1
style(ws, f'C{r}', '層2なし', F_BASE, border=True)
style(ws, f'D{r}', 'OSSが無い場合（運転員頼み）', F_BASE, border=True)
style(ws, f'E{r}', f'={F0}*{P1}', F_BASE, border=True, fmt=SCI)
style(ws, f'F{r}', f'=E{r}/E{base_row}', F_BASE, border=True, fmt='0.0')
r += 1
style(ws, f'C{r}', '層を1つ追加', F_BASE, border=True)
style(ws, f'D{r}', '独立したインターロックをもう1層追加（PFDは層2と同等と仮定）', F_BASE, border=True, wrap=True)
style(ws, f'E{r}', f'={F0}*{P1}*{P2}*{P2}', F_BASE, border=True, fmt=SCI)
style(ws, f'F{r}', f'=E{r}/E{base_row}', F_BASE, border=True, fmt='0.000')
r += 2
style(ws, f'C{r}', '（読み方）層を1つ追加すると頻度は追加層のPFD倍（約1/12）に下がる。'
      'ただし追加層が既存層と検出器やDCSを共有すると効果はここまで出ない（コモン事象）。', F_NOTE, wrap=True)
ws.row_dimensions[r].height = 30

# ============================================================ 6. リスクマトリックス
ws = wb.create_sheet('リスクマトリックス')
widths(ws, {'A': 3, 'B': 12, 'C': 22, 'D': 46, 'E': 3, 'F': 3, 'G': 10, 'H': 16, 'I': 16, 'J': 16, 'K': 16})
style(ws, 'B1', 'リスクマトリックス（仮置き・要確認）', F_TITLE)
style(ws, 'B2', 'この区分・ゾーン割付は一般形の仮置き。HAZchart解析基準（C-2-3）の正式な表と照合して確定すること。',
      F_WARN, wrap=True)

style(ws, 'B4', '頻度レベル（年間発生頻度）', F_SECTION, fill=FILL_SEC)
freq_rows = [
    ('E', '0.1回/年 以上', '年に何回も起こりうる'),
    ('D', '0.01以上 0.1未満', '検出が弱いと起こる'),
    ('C', '0.0001以上 0.01未満', '管理されていれば稀'),
    ('B', '0.000001以上 0.0001未満', '多重故障が必要'),
    ('A', '0.000001未満', '実質起こらない'),
]
r = 5
for lv, rng, desc in freq_rows:
    style(ws, f'B{r}', lv, F_BOLD, border=True, align='center')
    style(ws, f'C{r}', rng, F_BASE, border=True)
    style(ws, f'D{r}', desc, F_BASE, border=True)
    r += 1

style(ws, 'B12', '影響度レベル', F_SECTION, fill=FILL_SEC)
cons_rows = [
    ('1', '事故', '火災・爆発・大量漏洩・重大な人的被害'),
    ('2', '重大トラブル', '漏洩・機器損傷・長期の運転停止'),
    ('3', 'トラブル', '運転阻害・軽微な損傷（1週間未満で復旧）'),
    ('4', '小トラブル', '放出なし・プロセス乱れのみ'),
]
r = 13
for lv, name, desc in cons_rows:
    style(ws, f'B{r}', lv, F_BOLD, border=True, align='center')
    style(ws, f'C{r}', name, F_BASE, border=True)
    style(ws, f'D{r}', desc, F_BASE, border=True)
    r += 1

style(ws, 'G16', 'マトリックス（行=頻度レベル、列=影響度レベル）', F_SECTION, fill=FILL_SEC)
for i, h in enumerate(['頻度', '影響度1', '影響度2', '影響度3', '影響度4']):
    style(ws, f'{get_column_letter(7 + i)}17', h, F_BOLD, fill=FILL_HDR, border=True, align='center')
zones = [
    ('E', ['赤', '赤', '黄', '黄']),
    ('D', ['赤', '赤', '黄', '緑']),
    ('C', ['赤', '黄', '緑', '緑']),
    ('B', ['黄', '緑', '緑', '緑']),
    ('A', ['緑', '緑', '緑', '緑']),
]
zfill = {'赤': FILL_RED, '黄': FILL_YEL, '緑': FILL_GRN}
r = 18
for lv, zs in zones:
    style(ws, f'G{r}', lv, F_BOLD, border=True, align='center')
    for i, z in enumerate(zs):
        style(ws, f'{get_column_letter(8 + i)}{r}', z, F_BOLD, fill=zfill[z], border=True, align='center')
    r += 1
r += 1
style(ws, f'G{r}', '赤=許容不可（対策必須） 黄=条件付き（管理・監視強化のうえ猶予） 緑=許容可', F_NOTE)

for sheet in wb.worksheets:
    sheet.sheet_view.showGridLines = False

OUT = '/home/user/Base/outputs/HAZchart計算ベース_rev20260721.xlsx'
wb.save(OUT)
print('saved:', OUT)
