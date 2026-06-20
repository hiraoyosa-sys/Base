# -*- coding: utf-8 -*-
"""EG品質／DEG色相 解説 Word生成（読み物・配布用）。

ルール準拠：白地黒文字・モノクロ／本文に「→」記号書きを使わない／程度語・造語を使わない／
外部文献・AI痕跡を入れない／作成者メタは平尾名義。構成は 事実→仮説→対応。
"""
import os
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import eg_content as C

OUT = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "outputs")
os.makedirs(OUT, exist_ok=True)

FONT = "Meiryo"
INK = RGBColor(0x00, 0x00, 0x00)
SUB = RGBColor(0x40, 0x40, 0x40)

doc = Document()

# ---- 既定スタイル（フォント・東アジアフォント）----
normal = doc.styles["Normal"]
normal.font.name = FONT
normal.font.size = Pt(10.5)
normal.font.color.rgb = INK
normal.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
normal.paragraph_format.line_spacing = 1.15
normal.paragraph_format.space_after = Pt(4)

for sty in ["Heading 1", "Heading 2", "Heading 3", "Title"]:
    st = doc.styles[sty]
    st.font.name = FONT
    st.font.color.rgb = INK
    rpr = st.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), FONT)

# ---- ページ（A4縦・余白2cm）----
sec = doc.sections[0]
sec.page_width = Cm(21.0)
sec.page_height = Cm(29.7)
sec.left_margin = sec.right_margin = Cm(2.0)
sec.top_margin = sec.bottom_margin = Cm(2.0)


def _font(run, size=10.5, bold=False, color=INK):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    rpr = run._r.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts"); rpr.append(rf)
    rf.set(qn("w:eastAsia"), FONT)


def h1(text):
    p = doc.add_heading(level=1)
    r = p.add_run(text); _font(r, 15, True, INK)
    return p


def h2(text):
    p = doc.add_heading(level=2)
    r = p.add_run(text); _font(r, 12.5, True, INK)
    return p


def para(text, size=10.5, bold=False, color=INK, after=4):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    r = p.add_run(text); _font(r, size, bold, color)
    return p


def bullet(text, size=10.5):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(text); _font(r, size)
    return p


def num(text, size=10.5):
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(text); _font(r, size)
    return p


def shade(cell, hexc):
    tcPr = cell._tc.get_or_add_tcPr()
    sh = OxmlElement("w:shd")
    sh.set(qn("w:val"), "clear"); sh.set(qn("w:color"), "auto"); sh.set(qn("w:fill"), hexc)
    tcPr.append(sh)


def table(headers, rows, widths=None, head_size=9.5, body_size=9.5, aligns=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = t.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = hdr[i].paragraphs[0].add_run(h); _font(r, head_size, True, INK)
        shade(hdr[i], "D9D9D9")
    aligns = aligns or [WD_ALIGN_PARAGRAPH.LEFT] * len(headers)
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].paragraphs[0].alignment = aligns[i]
            r = cells[i].paragraphs[0].add_run(str(v)); _font(r, body_size)
    if widths:
        for i, w in enumerate(widths):
            for row in t.rows:
                row.cells[i].width = w
    # 段落間隔を詰める
    for row in t.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(1)
                p.paragraph_format.line_spacing = 1.05
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


LEFT = WD_ALIGN_PARAGRAPH.LEFT
CTR = WD_ALIGN_PARAGRAPH.CENTER

# ================= 表題 =================
tp = doc.add_paragraph()
tp.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = tp.add_run(C.TITLE); _font(r, 20, True, INK)
sp = doc.add_paragraph(); sp.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = sp.add_run(C.SUBTITLE); _font(r, 12, False, SUB)
dp = doc.add_paragraph(); dp.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = dp.add_run(f"{C.DATE}　プロセス技術部　EOGモデルチェンジPJ"); _font(r, 10, False, SUB)
doc.add_paragraph()

# ================= 1. はじめに =================
h1("1. はじめに（この資料の目的）")
for line in C.PURPOSE:
    para(line)

# ================= 2. 結論 =================
h1("2. 結論（要点）")
for line in C.SUMMARY:
    bullet(line)

# ================= 3. 現状（事実）=================
h1("3. 現状（事実）")
h2("3.1 製品別の品質状況")
para("製品ごとの品質状況と特徴的な事象を次に示す。色相悪化はタンク保管中に進み、採取直後のストリーム品はAPHA0であることが要点である。")
table(["製品／系", "状況", "特徴的な事象"],
      [[a, b, c2] for a, b, c2 in C.PRODUCT_STATUS],
      widths=[Cm(2.6), Cm(7.0), Cm(7.0)],
      aligns=[CTR, LEFT, LEFT])
para("顧客の対応窓口は、MEG（K値・HUV220）が東レ、TEG酸化指標が東洋紡、DEG色相がレゾナックである。")

h2("3.2 運転対応の実績")
para("色相とUVに対して試した運転操作と、その効きを次に示す。")
table(["操作", "期間", "HUV220低減", "DEG色相改善"],
      [list(r) for r in C.UNTEN_JISSEKI],
      widths=[Cm(7.2), Cm(3.2), Cm(3.0), Cm(3.0)],
      aligns=[LEFT, CTR, CTR, CTR])
para("要点：" + C.UNTEN_NOTE, bold=True)

h2("3.3 確認された事実")
para("これまでに確認された事実を整理する。特に、酸素雰囲気で色相悪化が止まる点と、脱水塔下流のリボイラーに穴が見つかった点は見立てに大きく効く。")
for k, v in C.FACTS:
    bullet(f"{k}　{v}")

# ================= 4. 着色の見立て（仮説）=================
h1("4. 着色の見立て（仮説）")
h2("4.1 工程に沿った筋")
para("原料アルデヒドの増加を起点に、前駆体が下流へ運ばれ、製品タンクの長期保管で色がつくと見ている。工程ごとに、起きていることと根拠を示す。")
table(["工程", "起きていること", "根拠"],
      [[a, b, c2] for a, b, c2 in C.SCENARIO],
      widths=[Cm(3.2), Cm(6.6), Cm(6.8)],
      aligns=[LEFT, LEFT, LEFT])
para("筋：" + C.SCENARIO_SUJI)

h2("4.2 重要な更新（6/17〜6/20）")
para("見立ての更新点を示す。なかでもリボイラー穴による水混入は、これまでの水分の整理を見直す必要がある。")
for k, v in C.UPDATES:
    bullet(f"{k}：{v}")

# ================= 5. 対応 =================
h1("5. 対応")
h2("5.1 薬剤（NaOHとリン酸Na）")
para("色相に効いたのはリン酸Na添加のみである。リン酸は継続投入が前提であり、止めると戻る。NaOHは有機酸の中和には効くが、塩基としてアルドール縮合を促す面がある。")
para("NaOH", bold=True)
for k in ["投入箇所", "投入量", "効果", "機構", "留意"]:
    bullet(f"{k}：{C.NAOH[k]}")
para("リン酸Na（Na₃PO₄）", bold=True)
bullet(f"投入箇所：{C.PHOS['投入箇所']}")
bullet(f"投入量：{C.PHOS['投入量']}")
bullet(f"効果：{C.PHOS['効果']}")
bullet(f"機構（DEG側）：{C.PHOS['機構DEG']}")
bullet(f"機構（MEG側）：{C.PHOS['機構MEG']}")
bullet(f"留意：{C.PHOS['留意']}")
para("リン酸Na テスト履歴", bold=True)
table(["回", "時期", "条件", "結果"],
      [list(r) for r in C.PHOS_TEST],
      widths=[Cm(1.2), Cm(3.0), Cm(5.5), Cm(7.0)],
      aligns=[CTR, LEFT, LEFT, LEFT])

h2("5.2 対応方案（方案1〜7）")
para(C.HOUAN_HOSHIN)
table(["方案", "内容", "戦略", "ねらい", "現状"],
      [list(r) for r in C.HOUAN],
      widths=[Cm(1.6), Cm(3.8), Cm(2.4), Cm(4.0), Cm(4.6)],
      head_size=9, body_size=8.5,
      aligns=[CTR, LEFT, CTR, LEFT, LEFT])

h2("5.3 スタートアップ運転条件案")
para(C.SU_NOTE)
table(["項目", "方針"],
      [[a, b] for a, b in C.SU_JOKEN],
      widths=[Cm(4.0), Cm(12.6)],
      aligns=[LEFT, LEFT])

h2("5.4 早期色相判断・指標")
table(["指標", "内容"],
      [[a, b] for a, b in C.SHIHYO],
      widths=[Cm(3.6), Cm(13.0)],
      aligns=[LEFT, LEFT])

h2("5.5 異物の確認と顧客への説明")
para("異物・タンク・ライン処理", bold=True)
for a, b in C.IBUTSU:
    bullet(f"{a}：{b}")
para("顧客（レゾナック）向け説明の骨子", bold=True)
for line in C.KOKYAKU:
    bullet(line)
para(C.KOKYAKU_NOTE, color=SUB)

# ================= 6. M50との接続 =================
h1("6. M50（モデルチェンジ）との接続")
for line in C.M50:
    para(line)

# ================= 7. 次アクションと要検証 =================
h1("7. 次アクションと要検証")
h2("7.1 次アクション")
for v in C.NEXT:
    num(v)
h2("7.2 要検証・齟齬")
for v in C.OPEN:
    bullet(v)

# ---- メタデータ（作成者：平尾名義／AI痕跡なし）----
cp = doc.core_properties
cp.author = C.AUTHOR
cp.last_modified_by = C.AUTHOR
cp.title = "EG品質の現状とDEG色相悪化への対応"
cp.comments = ""

path = os.path.join(OUT, "EG品質_DEG色相_解説.docx")
doc.save(path)
print("saved:", path)
