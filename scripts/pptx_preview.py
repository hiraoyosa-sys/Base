# -*- coding: utf-8 -*-
"""pptx簡易プレビュー：図形座標・塗り・テキスト・埋め込み画像・表をmatplotlibで再現してPNG化。
体裁（重なり・密度・色）の目視確認用。完全な再現ではないがレイアウト確認に十分。"""
import os, io, sys
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.patches import FancyBboxPatch, Rectangle
import matplotlib.font_manager as fm
from PIL import Image
from pptx import Presentation
from pptx.util import Emu
from pptx.enum.shapes import MSO_SHAPE_TYPE

fp = fm.FontProperties(fname="/usr/share/fonts/opentype/ipafont-gothic/ipag.ttf")
EMU = 914400.0
SW, SH = 13.333, 7.5

def emu_in(v):
    try: return v / EMU
    except Exception: return 0

def rgb_of(color):
    try:
        c = color.rgb
        return "#%02X%02X%02X" % (c[0], c[1], c[2])
    except Exception:
        return None

def shape_fill(sh):
    try:
        if sh.fill.type is not None:
            return rgb_of(sh.fill.fore_color)
    except Exception:
        pass
    return None

def shape_line(sh):
    try:
        return rgb_of(sh.line.color)
    except Exception:
        return None

def draw_text_frame(ax, tf, x, y, w, h):
    # paragraphs stacked top-down
    cy = y + h - 0.08
    for p in tf.paragraphs:
        runs = p.runs
        if not runs:
            cy -= 0.18; continue
        txt = "".join(r.text for r in runs)
        if not txt.strip():
            cy -= 0.12; continue
        r0 = runs[0]
        try: sz = r0.font.size.pt if r0.font.size else 12
        except Exception: sz = 12
        col = rgb_of(r0.font.color) or "#222222"
        bold = bool(r0.font.bold)
        # wrap approx by width
        maxchars = max(8, int(w / (sz * 0.011)))
        line = txt if len(txt) <= maxchars else txt[:maxchars-1] + "…"
        ax.text(x + 0.1, cy, line, ha="left", va="top", fontproperties=fp,
                fontsize=min(sz, 20) * 0.9, color=col, fontweight=("bold" if bold else "normal"))
        cy -= (sz * 0.020 + 0.08)
        if cy < y: break

def draw_table(ax, tbl, x, y, w, h):
    rows = len(tbl.rows); cols = len(tbl.columns)
    cw = [emu_in(c.width) for c in tbl.columns]
    tot = sum(cw) or w
    cw = [c / tot * w for c in cw]
    rh = h / rows
    yy = y + h
    for i in range(rows):
        yy -= rh; xx = x
        for j in range(cols):
            cell = tbl.cell(i, j)
            fillc = None
            try: fillc = rgb_of(cell.fill.fore_color)
            except Exception: pass
            ax.add_patch(Rectangle((xx, yy), cw[j], rh, facecolor=(fillc or "white"),
                                   edgecolor="#C9D2D8", linewidth=0.5))
            t = cell.text.replace("\n", " ")
            if t.strip():
                mc = max(6, int(cw[j] / 0.07))
                ax.text(xx + 0.03, yy + rh - 0.05, t if len(t) <= mc else t[:mc-1] + "…",
                        ha="left", va="top", fontproperties=fp, fontsize=6.2, color="#222222")
            xx += cw[j]

def render(path, outdir):
    prs = Presentation(path)
    os.makedirs(outdir, exist_ok=True)
    outs = []
    for idx, slide in enumerate(prs.slides, 1):
        fig, ax = plt.subplots(figsize=(SW, SH), dpi=70)
        ax.set_xlim(0, SW); ax.set_ylim(0, SH); ax.axis("off")
        ax.add_patch(Rectangle((0, 0), SW, SH, facecolor="white", edgecolor="#999999", linewidth=1))
        for sh in slide.shapes:
            x, y0, w, h = emu_in(sh.left), emu_in(sh.top), emu_in(sh.width), emu_in(sh.height)
            y = SH - y0 - h  # invert
            st = sh.shape_type
            if st == MSO_SHAPE_TYPE.PICTURE:
                try:
                    img = Image.open(io.BytesIO(sh.image.blob))
                    ax.imshow(img, extent=(x, x + w, y, y + h), aspect="auto", zorder=5)
                except Exception:
                    ax.add_patch(Rectangle((x, y), w, h, facecolor="#EEEEEE", edgecolor="#AAAAAA"))
                continue
            if sh.has_table:
                draw_table(ax, sh.table, x, y, w, h); continue
            # autoshape / textbox
            fillc = shape_fill(sh); linec = shape_line(sh)
            if st == MSO_SHAPE_TYPE.AUTO_SHAPE or fillc:
                ax.add_patch(FancyBboxPatch((x, y), w, h, boxstyle="round,pad=0.01,rounding_size=0.05",
                             facecolor=(fillc or "none"), edgecolor=(linec or "none"),
                             linewidth=1.0 if linec else 0))
            if sh.has_text_frame and sh.text_frame.text.strip():
                draw_text_frame(ax, sh.text_frame, x, y, w, h)
        plt.subplots_adjust(left=0, right=1, top=1, bottom=0)
        op = os.path.join(outdir, f"slide{idx:02d}.png")
        plt.savefig(op, dpi=70); plt.close(fig); outs.append(op)
    return outs

if __name__ == "__main__":
    src = sys.argv[1] if len(sys.argv) > 1 else os.path.join(os.path.dirname(__file__), "..", "outputs", "製品DEG色相悪化_対策会議.pptx")
    outdir = sys.argv[2] if len(sys.argv) > 2 else "/tmp/preview"
    outs = render(src, outdir)
    print("rendered", len(outs), "slides ->", outdir)
