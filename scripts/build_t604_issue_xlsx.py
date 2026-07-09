# -*- coding: utf-8 -*-
"""T-604 DEG色相悪化：イシュー特定と課題の分離（たたき台）Excel生成。

シート1＝課題の分離表（縦＝分離軸、横＝仮説・成立条件・既存整理・開放所見・確認データ）。
シート2＝動かない事実・イシューの構図・情報待ちリスト。
配色はモノクロ（白地・黒文字・細罫線・見出しは薄グレー）＝DEG着色_フロー整理.xlsxと同体裁。
"""
import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

OUT = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "outputs")
os.makedirs(OUT, exist_ok=True)

DATE = "2026-07-09"
TITLE = "T-604 DEG色相悪化　イシュー特定と課題の分離（たたき台）"

HEADER_FILL = "E6E6E6"
LABEL_FILL = "F2F2F2"
SUB_FILL = "F7F7F7"
WHITE = "FFFFFF"
INK = "000000"
SUB_INK = "404040"
BORDER_CLR = "808080"

FONT = "IPAGothic"
thin = Side(style="thin", color=BORDER_CLR)
border = Border(left=thin, right=thin, top=thin, bottom=thin)
med = Side(style="medium", color="595959")


def fill(c):
    return PatternFill("solid", fgColor=c)


def est_lines(text, cpl):
    n = 0
    for para in text.split("\n"):
        n += max(1, -(-len(para) // cpl))
    return n


# ===== シート1：課題の分離表 =====
# 列定義（たたき台。列の意味は木村さんの用途に合わせて要調整）
COL_DEFS = [
    ("分離軸\n（T-604と他タンクの\n違いの候補）", 17),
    ("仮説\n（この違いがあればT-604だけの悪化を説明できる）", 36),
    ("妥当と言えるための条件\n（何が確認できれば成り立つか）", 36),
    ("既存整理での位置づけ", 32),
    ("タンク開放所見との関係", 30),
    ("確認に必要なデータ・分析", 32),
]

ROWS = [
    (
        "液の中身\n（受入・ブレンド）",
        "悪化品（ポリエナール前駆体を含むDEG）の受入がT-604に集中し、前駆体濃度が最も高い液がここで長期滞留した。T-555・T-615にはブレンド・希釈後の液が回るため悪化が目立たない。",
        "受払履歴で、悪化期間の初期流動品・悪化品の入槽先がT-604に偏っていたこと。T-555・T-615への移送がブレンド後だったこと。",
        "既存メカニズム（450nm成長の場はタンク長滞留・前駆体が反応しきるとAPHA30前後で安定）と最も整合。初期タンク運用「T-604で受け、T-555・T-615へ展開しブレンドで調整」とも符合。",
        "前駆体・重質物が最も多く通過・滞留したタンクなら、内面の付着・変色が最も強く出ることも同時に説明できる（付着は悪化の結果、という向きと同じ）。",
        "タンク受払履歴（悪化期の受入元・量・ブレンド比率）。残液・付着物のUVチャートとGPC（T-555残液との比較）。",
    ),
    (
        "塩基性\n（Na・pH）",
        "T-604内のNa濃度・pHが他タンクより高く、「A-ALDとNaの併存」の悪化条件がタンク内でも成立している。",
        "T-604残液のNa・pHがT-555・T-615より有意に高いこと。Na投入管理履歴と品質の相関で、Na添加期の液がT-604に偏って入っていたこと。",
        "主因の現行整理（A-ALD濃度上昇とNa存在の組み合わせ・7/1ほぼ合意）と直結。「Na停止でストリームは改善する一方、特定タンク(604)で濃度上昇事例あり」という既往記録とも符合。",
        "塩基性そのものは内面の腐食・汚れを直接は説明しない（腐食は有機酸側の論点として分離）。",
        "残液のpH・Na実測。Na投入・リン酸Na転換の時期とT-604受入時期の重なり。2008年比較はNa無添加期のため単純比較に限界がある点に留意。",
    ),
    (
        "滞留時間\n・残液運用",
        "T-604は回転が遅く残液が長く残る運用のため、鎖伸長（脱水縮合）の時間が他タンクより長い。",
        "受払データから各タンクの平均滞留日数・残液量を出し、T-604が突出していること。",
        "「450nm成長の場はタンク（長滞留）」と整合。本管DEGは27日でAPHA5以下・タンク品で450nmが成長する、という事実とも同じ向き。",
        "長滞留であればスラッジ・付着の蓄積も説明しやすい。",
        "タンク別の滞留日数・残液量の試算（受払履歴から）。",
    ),
    (
        "タンク内面の状態\n（腐食・付着物）",
        "T-604内面の錆・付着物が着色の原因側（触媒または供給源）として働いている。",
        "T-604付着物に、450nm成長へ直接つながる触媒作用または特異成分が確認されること。既往試験（鉄錆単独では着色しない・C-1503付着物はピーク380nmで実機450nmと別現象・黒色付着物混入でも450nmピーク再現せず）を覆す結果が必要。",
        "既存整理では否定側が優勢（鉄サビは450nmの主因でない・サビは原因でなく結果の可能性＝前田）。ただし局所腐食・380nm寄与までは無罪と言い切っていない＝未確定を残す。",
        "開放所見「状態が良くない」（詳細は情報待ち）を直接説明する仮説。ただし逆向きの説明（アルデヒド高濃度で生成した有機酸（ギ酸）がタンクを腐食させた＝悪化の結果）も既に立っており、原因か結果かの向き決めが本丸の論点。",
        "付着物の元素・有機分析（Fe・P・Na・有機成分。T-555異物「ほぼ鉄」との比較）。付着物のDEG共存保管試験（450nm再現の有無）。肉厚測定と腐食部位の分布（液面部・底部・気相部のどこか）。",
    ),
    (
        "雰囲気\n（窒素シール・酸素）",
        "T-604は酸素が特に入らず、悪化を止める側に働く酸素の効果が効いていない。または逆に微量のエアリークがあり、活性金属の生成に働いている（木村コメント＝活性金属の発生に微量酸素が必要）。",
        "タンク別の窒素シール・呼吸・エアリークの実態に、T-604だけの差があること。",
        "酸素雰囲気で悪化が止まる事実はあるが、T-555も窒素シールで残液はAPHA30で安定しており、雰囲気の差だけでは説明力が弱い。窒素抑制の機序自体が3説併存（ラジカル停止・CO2溶解pH低下・吸湿水分）で未確定。",
        "エアリークや結露があれば、気相部の内面腐食と符合する。",
        "シール窒素の設定・パージ量・酸素濃度の実測。腐食部位の分布との突き合わせ。",
    ),
    (
        "水分・温度",
        "T-604の液水分が特に低く鎖伸長が進みやすい。または温度が高く反応が速い。",
        "タンク別の水分・温度実績に有意差があること。",
        "水分の寄与度は感度2〜3割で未確定。温度はDEG色相への検証が弱いとして優先度を下げた経緯（木村判断・6/3）。",
        "直接は説明しない。",
        "タンク別の水分・温度データの比較（あれば）。",
    ),
    (
        "見え方\n（監視・運用の偏り）",
        "タンク固有の差ではなく、悪化品の投入とUV450監視がT-604に集中しているため「T-604だけ悪化」に見えている。",
        "同等の悪化品をT-555・T-615に入れた場合に同程度の速度（1日あたり0.004程度）でUV450が上昇すれば、タンク固有差は否定される。",
        "T-604はUV450早期判定の実績タンク（4/13からのデータで悪化品投入時1日あたり約0.004上昇・正常品はベース0.002から動かない）＝監視が最も密なタンクであることは事実。",
        "開放所見は説明しない（設備側の課題として分離）。",
        "入槽先別のUV450推移の比較。製品ストリームのボンベサンプリング（色相悪化がプラント由来かタンク由来かの切り分け・向後提案）。",
    ),
]

# ===== シート2：イシューの構図・動かない事実・情報待ち =====
ISSUES = [
    ("イシュー1\n（液の課題）",
     "T-604だけDEG色相悪化が進むのはなぜか。主仮説は「何が入ったか（前駆体）」「何と一緒か（Na）」「どれだけ置いたか（滞留）」の組み合わせ＝タンクの個体差ではなく運用の差で説明する筋が、既存整理と最も整合する。"),
    ("イシュー2\n（設備の課題）",
     "T-604開放で内面の状態が良くないのはなぜか・健全性をどう回復するか。既存整理では「腐食・付着は悪化の結果（アルデヒド高濃度で生成した有機酸による腐食）」の可能性が高い側。色相の原因究明とは切り離し、設備健全性（肉厚・補修）として進める。"),
    ("両者の接続点",
     "付着物分析だけが両イシューをつなぐ（付着・錆が色相悪化の原因か、悪化の結果かの向き決め）。ここを最初に確定させると課題が完全に分離できる。"),
]

FACTS = [
    ("1", "着色物質は共役ポリエナール類（基礎研GPCで実測・MW200〜800）。主因はA-ALD由来（F-ALD・G-ALDも一部寄与）。"),
    ("2", "450nm成長の場はタンク（長滞留）。前駆体が反応しきるとAPHA30前後で安定（T-555・T-604残液と整合）。"),
    ("3", "酸素雰囲気（船・ローリー）では悪化が止まる。鉄錆単独では着色しない（保管試験の実測）。"),
    ("4", "運用で色相に効いたのはリン酸Na（Na3PO4）のみ。継続投入が前提（止めると戻る）。"),
    ("5", "主要因は「A-ALD濃度上昇」と「Naの存在」の組み合わせでほぼ合意（7/1）。2008年にアセト高でも問題が出なかったのはNa無添加のため。"),
    ("6", "T-604のUV450実績＝悪化品投入時は1日あたり約0.004上昇・正常品はベース0.002から動かない（4/13以降）。"),
    ("7", "T-555（6/16初回開放）は黒色異物少なく比較的綺麗・異物はほぼ鉄。T-615（6/17）もT-555と同程度。T-555残液は5/18以降APHA上昇なし。"),
    ("8", "今回のT-604開放では内面の状態が良くない（所見の詳細は情報収集中）。"),
]

PENDING = [
    "T-604開放所見の詳細（部位・色・付着量・腐食の程度・写真）",
    "タンク受払履歴（悪化期の受入元・ブレンド比率・滞留日数）",
    "残液・付着物の分析結果（UVチャート・pH・Na・Fe・P）",
    "窒素シール・酸素濃度の実態（タンク別）",
]

NOTE = ("本表はたたき台。列の意味（特に「確認に必要なデータ・分析」が何に向けた確認か）と行の粒度は、"
        "木村さんの用途に合わせて調整する。具体データが入り次第、各行の「妥当と言えるための条件」を満たすか順に潰していく。")


wb = Workbook()
ws = wb.active
ws.title = "課題の分離"
ws.sheet_view.showGridLines = False

NCOL = len(COL_DEFS)
for i, (_, w) in enumerate(COL_DEFS):
    ws.column_dimensions[get_column_letter(1 + i)].width = w
last_letter = get_column_letter(NCOL)
CPLS = [int(w / 2.05) for _, w in COL_DEFS]

r = 1
ws.merge_cells(f"A{r}:{last_letter}{r}")
c = ws.cell(r, 1, TITLE)
c.font = Font(name=FONT, size=15, bold=True, color=INK)
c.alignment = Alignment(horizontal="left", vertical="center")
ws.row_dimensions[r].height = 26
r += 1
ws.merge_cells(f"A{r}:{last_letter}{r}")
c = ws.cell(r, 1, f"（{DATE} 時点・たたき台）")
c.font = Font(name=FONT, size=9, color=SUB_INK)
c.alignment = Alignment(horizontal="right", vertical="center")
for cidx in range(1, NCOL + 1):
    ws.cell(r, cidx).border = Border(bottom=med)
ws.row_dimensions[r].height = 14
r += 1

HEADER_ROW = r
for i, (name, _) in enumerate(COL_DEFS):
    cell = ws.cell(r, 1 + i, name)
    cell.font = Font(name=FONT, size=9.5, bold=True, color=INK)
    cell.fill = fill(HEADER_FILL)
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border = Border(left=thin, right=thin, top=thin, bottom=med)
ws.row_dimensions[r].height = 44
r += 1

for row in ROWS:
    maxlines = 1
    for i, txt in enumerate(row):
        cell = ws.cell(r, 1 + i, txt)
        if i == 0:
            cell.font = Font(name=FONT, size=9, bold=True, color=INK)
            cell.fill = fill(LABEL_FILL)
            cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True, indent=1)
        else:
            cell.font = Font(name=FONT, size=8.5, color=INK if i in (1, 2) else SUB_INK)
            cell.fill = fill(WHITE)
            cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True, indent=1)
        cell.border = border
        maxlines = max(maxlines, est_lines(txt, CPLS[i]))
    ws.row_dimensions[r].height = maxlines * 12.5 + 7
    r += 1

r += 1
ws.merge_cells(f"A{r}:{last_letter}{r}")
c = ws.cell(r, 1, NOTE)
c.font = Font(name=FONT, size=9, color=SUB_INK)
c.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True, indent=1)
for cidx in range(1, NCOL + 1):
    ws.cell(r, cidx).border = border
ws.row_dimensions[r].height = est_lines(NOTE, sum(CPLS)) * 15 + 10

ws.print_title_rows = f"{HEADER_ROW}:{HEADER_ROW}"
ws.page_setup.orientation = "landscape"
ws.page_setup.fitToWidth = 1
ws.page_setup.fitToHeight = 0
ws.sheet_properties.pageSetUpPr.fitToPage = True
ws.page_margins.left = ws.page_margins.right = 0.3
ws.page_margins.top = ws.page_margins.bottom = 0.4
ws.freeze_panes = "B" + str(HEADER_ROW + 1)


ws2 = wb.create_sheet("事実と情報待ち")
ws2.sheet_view.showGridLines = False
ws2.column_dimensions["A"].width = 16
ws2.column_dimensions["B"].width = 104


def s2_band(r, text):
    ws2.merge_cells(f"A{r}:B{r}")
    c = ws2.cell(r, 1, text)
    c.font = Font(name=FONT, size=12, bold=True, color=INK)
    c.fill = fill(HEADER_FILL)
    c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    c.border = border
    ws2.cell(r, 2).border = border
    ws2.row_dimensions[r].height = 22
    return r + 1


def s2_kv(r, k, v):
    kc = ws2.cell(r, 1, k)
    kc.font = Font(name=FONT, size=9.5, bold=True, color=INK)
    kc.fill = fill(LABEL_FILL)
    kc.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    kc.border = border
    vc = ws2.cell(r, 2, v)
    vc.font = Font(name=FONT, size=9.5, color=INK)
    vc.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True, indent=1)
    vc.border = border
    ws2.row_dimensions[r].height = max(26, est_lines(v, 52) * 15 + 7)
    return r + 1


r = 1
r = s2_band(r, "イシューの構図（2つに分離する）")
for k, v in ISSUES:
    r = s2_kv(r, k, v)
r += 1
r = s2_band(r, "動かない事実（前提）")
for k, v in FACTS:
    r = s2_kv(r, k, v)
r += 1
r = s2_band(r, "情報待ち（入り次第この表に反映）")
for i, v in enumerate(PENDING, 1):
    r = s2_kv(r, str(i), v)

ws2.page_setup.orientation = "landscape"
ws2.page_setup.fitToWidth = 1
ws2.page_setup.fitToHeight = 0
ws2.sheet_properties.pageSetUpPr.fitToPage = True

wb.properties.creator = "hirao kazuaki/0465811/平尾　一陽"
wb.properties.lastModifiedBy = "hirao kazuaki/0465811/平尾　一陽"
wb.properties.title = TITLE
wb.properties.description = ""
wb.properties.keywords = ""

path = os.path.join(OUT, "T-604_DEG色相_課題分離.xlsx")
wb.save(path)
print("saved:", path)
