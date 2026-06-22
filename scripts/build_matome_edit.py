# -*- coding: utf-8 -*-
"""まとめ資料(65f0890b)を直接修正：S6設備/S11基礎研の空欄記入、
『どこで・どういう反応』統合スライド追加、補足索引/用意資料スライド追加。新規作成でなく当該ファイルを編集。"""
import os
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, PP_PLACEHOLDER

SRC = "/root/.claude/uploads/0ac036e2-7cd8-5317-a99c-aa39021935be/65f0890b-__DEG________EO__.pptx"
OUT = os.path.join(os.path.dirname(__file__), "..", "outputs", "DEG色相対応_SU前会議.pptx")
IN = 914400
EA = "游ゴシック"
BLUE = RGBColor(0x00, 0x5B, 0xAB); DBLUE = RGBColor(0x00, 0x3F, 0x7E); BLACK = RGBColor(0x00, 0x00, 0x00)
INK = RGBColor(0x22, 0x22, 0x22); GRAY = RGBColor(0x66, 0x66, 0x66); WHITE = RGBColor(0xFF, 0xFF, 0xFF)
H_FILL = RGBColor(0xEB, 0xEF, 0xF2); PROC = RGBColor(0xD7, 0xE0, 0xE5); AMBER = RGBColor(0xFF, 0xF6, 0xE0)
AMBER_T = RGBColor(0xC0, 0x7A, 0x00)
LINE = RGBColor(0xC9, 0xD2, 0xD8)

prs = Presentation(SRC)
W = prs.slide_width / IN

def _ea(run):
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:latin", "a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None: el = rPr.makeelement(qn(tag), {}); rPr.append(el)
        el.set("typeface", EA)

def tbx(s, l, t, w, h):
    x = s.shapes.add_textbox(Emu(int(l*IN)), Emu(int(t*IN)), Emu(int(w*IN)), Emu(int(h*IN)))
    x.text_frame.word_wrap = True; return x.text_frame

def P(tf, text, size, bold=False, first=False, color=INK, bullet=False, space=5, align=None):
    p = tf.paragraphs[0] if (first and not tf.paragraphs[0].runs) else tf.add_paragraph()
    if align is not None: p.alignment = align
    r = p.add_run(); r.text = ("・" + text) if bullet else text
    r.font.size = Pt(size); r.font.bold = bold; r.font.color.rgb = color; _ea(r)
    p.space_after = Pt(space); return p

def rrect(s, l, t, w, h, fill, line=None, lw=1.0):
    sp = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Emu(int(l*IN)), Emu(int(t*IN)), Emu(int(w*IN)), Emu(int(h*IN)))
    sp.fill.solid(); sp.fill.fore_color.rgb = fill
    if line is None: sp.line.fill.background()
    else: sp.line.color.rgb = line; sp.line.width = Pt(lw)
    sp.shadow.inherit = False; return sp

def arrow(s, l, t, w, h, color=BLUE):
    sp = s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Emu(int(l*IN)), Emu(int(t*IN)), Emu(int(w*IN)), Emu(int(h*IN)))
    sp.fill.solid(); sp.fill.fore_color.rgb = color; sp.line.fill.background(); sp.shadow.inherit = False; return sp

def vbox(s, l, t, w, h, lines, fill, line=BLUE, lw=1.1):
    sp = rrect(s, l, t, w, h, fill, line=line, lw=lw)
    tf = sp.text_frame; tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = Emu(int(0.05*IN)); tf.margin_right = Emu(int(0.05*IN)); tf.margin_top = Emu(int(0.03*IN)); tf.margin_bottom = Emu(int(0.03*IN))
    for i, (t2, sz, b, c) in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph(); p.alignment = PP_ALIGN.CENTER
        r = p.add_run(); r.text = t2; r.font.size = Pt(sz); r.font.bold = b; r.font.color.rgb = c; _ea(r)
    return sp

def title_of(s):
    for ph in s.placeholders:
        if ph.placeholder_format.idx == 0: return ph.text_frame.text.strip()
    return ""

def slide_by_title(key):
    for s in prs.slides:
        if key in title_of(s): return s
    return None

def body_textbox(s, l, t, w, h, bullets, size=14):
    tf = tbx(s, l, t, w, h)
    for i, b in enumerate(bullets):
        if isinstance(b, tuple):
            txt, bold, col = b
            P(tf, txt, size, bold=bold, color=col, first=(i == 0), bullet=False, space=8)
        else:
            P(tf, b, size, bullet=True, first=(i == 0), space=8)
    return tf

# ── S6 設備開放で確認された事実（空欄を記入） ──
s6 = slide_by_title("設備開放で確認された事実")
if s6:
    body_textbox(s6, 0.7, 1.55, 12.0, 5.4, [
        ("■ 色相悪化物質ができる反応は、水が少なくなる脱水塔以降で起こると考えられる。そこでEG精製系（脱水塔以降の塔・配管・タンク）を中心に開放・検査した。", True, DBLUE),
        "開けた結果、DEG精製塔（C-1503）トップの黒色付着物（ほぼ鉄サビ）以外に、目立った汚れ・腐食・異物は見つからなかった。",
        "その黒色付着物（鉄サビ）をDEGに加えて保管したところ、色相悪化の450nmは現れなかった（基礎研）。＝鉄サビは色相悪化の原因ではない。",
        "各機器の肉厚を定修前に測定したが、有意な減りはなかった。",
        ("→ 設備や金属が原因で色相が悪くなったのではないと判断。残る上流側は色相悪化物質ができる反応の場ではないため、これ以上開ける必要はない。", True, DBLUE),
    ], size=14)

# ── S11 基礎研で検証された事実（空欄を記入） ──
s11 = slide_by_title("基礎研で検証された事実")
if s11:
    body_textbox(s11, 0.7, 1.5, 12.0, 5.6, [
        ("■ 色相悪化物質の正体を調べた", True, DBLUE),
        "着色物質をGPC・GCで分析した → 分子量200〜800の共役の長い化合物（ポリエナール類）と分かった。量はごく微量（ppbオーダー）。",
        ("■ 鉄サビで色相悪化が起きるかを調べた", True, DBLUE),
        "鉄サビ（C-1503付着物・タンク異物）をDEGに加えて保管した → 色相悪化の450nmは現れなかった。＝鉄サビは色相悪化物質を作らない。",
        ("■ アルデヒドだけで色相悪化が起きるかを調べた", True, DBLUE),
        "アルデヒド（A-ALD・F-ALD）をDEGに加えて保管した（12日） → 450nmは現れなかった。＝アルデヒドを加えただけでは色相悪化は再現しない。",
        "ラボでEG反応系を模擬してアルデヒドを加えた → UV（短波長側）が悪化した（F-ALDとA-ALDの両方があるとより大きい）。＝アルデヒドがUV悪化のもとになることは確認。",
        ("※ 色相悪化（450nm）そのものを再現するには至っていない（要追検証）。", False, GRAY),
    ], size=13)

# ── S14 運転対応実績：タイトル改名＋『DEG色相への影響』列を記入（たたき台） ──
s14 = slide_by_title("運転対応実績")
if s14:
    for ph in s14.placeholders:
        if ph.placeholder_format.idx == 0:
            tf = ph.text_frame; p0 = tf.paragraphs[0]
            if p0.runs:
                p0.runs[0].text = "運転対応実績とDEG色相への影響"
                for ex in p0.runs[1:]: ex._r.getparent().remove(ex._r)
            for p in tf.paragraphs[1:]: p._p.getparent().remove(p._p)
    tbl = None
    for sh in s14.shapes:
        if sh.has_table: tbl = sh.table
    if tbl is not None:
        tbl.columns[0].width = Emu(int(4.6*IN)); tbl.columns[1].width = Emu(int(1.9*IN)); tbl.columns[2].width = Emu(int(5.37*IN))
        WORSE = RGBColor(0xFC, 0xE4, 0xE4); WORSE_T = RGBColor(0xB0, 0x00, 0x00)
        NONEF = RGBColor(0xF4, 0xF6, 0xF8); NONET = RGBColor(0x55, 0x55, 0x55)
        def set_keep(cell, text):
            tf = cell.text_frame; p0 = tf.paragraphs[0]
            if p0.runs:
                p0.runs[0].text = text
                for ex in p0.runs[1:]: ex._r.getparent().remove(ex._r)
            else:
                r = p0.add_run(); r.text = text; _ea(r)
            for p in tf.paragraphs[1:]: p._p.getparent().remove(p._p)
        def fill_inf(cell, text, fill, color, bold=False, size=10):
            tf = cell.text_frame; tf.word_wrap = True; cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.margin_left = Emu(int(0.06*IN)); cell.margin_right = Emu(int(0.06*IN))
            p0 = tf.paragraphs[0]
            for ex in list(p0.runs): ex._r.getparent().remove(ex._r)
            for p in tf.paragraphs[1:]: p._p.getparent().remove(p._p)
            r = p0.add_run(); r.text = text; r.font.size = Pt(size); r.font.bold = bold; r.font.color.rgb = color; _ea(r)
            cell.fill.solid(); cell.fill.fore_color.rgb = fill
        set_keep(tbl.cell(0, 2), "DEG色相への影響")
        infmap = [
            ("NaOH", "悪化させた可能性：塩基がアルデヒドどうしの反応を後押し。→S/U時はEOの不純物低減を確認のうえ投入を判断", WORSE, WORSE_T, False),
            ("C-1501塔底温度", "悪化させた可能性：分解されず重い形で下流(DEG)へ回る（HUVは改善）。→S/U時は不純物・HUVを見て安易に下げない", WORSE, WORSE_T, False),
            ("Na3PO4", "改善傾向の実績あり（色相の進みが緩やかになる）", AMBER, AMBER_T, True),
        ]
        for i in range(1, len(tbl.rows)):
            item = tbl.cell(i, 0).text
            chosen = next(((t, f, c, b) for k, t, f, c, b in infmap if k in item), None)
            if chosen is None: chosen = ("明確な影響は確認されていない", NONEF, NONET, False)
            fill_inf(tbl.cell(i, 2), *chosen)
        lg = tbx(s14, 0.4, 6.45, 11.87, 0.5)
        P(lg, "※ DEG色相への影響は現時点の見立て（たたき台）。 赤＝悪化の可能性 ／ 橙＝改善傾向 ／ 無印＝明確な影響は確認されていない（HUV・ALDに有意変化なし）", 9.5, color=GRAY, first=True)

# ── 統合スライド『どこで・どういう反応が起こるか』（スライド5,6統合）を追加 ──
ly = None
for l in prs.slide_layouts:
    if l.name == "タイトルとコンテンツ": ly = l
if ly is None: ly = prs.slide_layouts[2]
s = prs.slides.add_slide(ly)
for ph in list(s.placeholders):
    if ph.placeholder_format.idx != 0 and ph.placeholder_format.type in (
        PP_PLACEHOLDER.BODY, PP_PLACEHOLDER.OBJECT, PP_PLACEHOLDER.SUBTITLE, PP_PLACEHOLDER.CENTER_TITLE):
        ph._element.getparent().remove(ph._element)
for ph in s.placeholders:
    if ph.placeholder_format.idx == 0:
        ph.text = ""; r = ph.text_frame.paragraphs[0].add_run(); r.text = "どこで色相悪化物質ができるか（工程と反応場）"
        r.font.size = Pt(24); r.font.bold = True; r.font.color.rgb = BLACK; _ea(r)
AMBER_T = RGBColor(0xC0, 0x7A, 0x00); GUIDE = RGBColor(0xB8, 0xC4, 0xCC)
# 上段：工程フロー（7段）。各工程の左端・中心を記録し、反応はその真下へ揃える
stages = ["EO反応系", "EG反応系", "EG濃縮系", "EG脱水系", "MEG精製系", "DEG精製系", "製品タンク"]
n = len(stages); bw = 1.55; gap = (W - 0.6 - n*bw) / (n - 1); x = 0.3; ytop = 1.45; bh = 0.78
lefts = []; cxs = []
for k, nm in enumerate(stages):
    last = (k == n-1)
    vbox(s, x, ytop, bw, bh, [(nm, 11, True, (AMBER_T if last else DBLUE))],
         (AMBER if last else H_FILL), line=(AMBER_T if last else BLUE))
    lefts.append(x); cxs.append(x + bw/2); x += bw + gap
for k in range(n-1):
    arrow(s, lefts[k]+bw+0.01, ytop+bh/2-0.08, gap-0.02, 0.16)

# 帯：どこが「反応場」か（工程の真下に色分け）
zy = 2.42; zh = 0.36
zb1 = rrect(s, lefts[0], zy, (lefts[2]+bw)-lefts[0], zh, RGBColor(0xF0, 0xF2, 0xF4), line=LINE, lw=0.75)
zb1.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
p = zb1.text_frame.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
r = p.add_run(); r.text = "水が多い → アルデヒドはパージで系外（まだ色はつかない）"; r.font.size = Pt(10.5); r.font.bold = True; r.font.color.rgb = GRAY; _ea(r)
zb2 = rrect(s, lefts[3], zy, (lefts[6]+bw)-lefts[3], zh, AMBER, line=AMBER_T, lw=1.0)
zb2.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
p = zb2.text_frame.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
r = p.add_run(); r.text = "色相悪化物質ができる反応場（水が少ない＝脱水塔以降）"; r.font.size = Pt(10.5); r.font.bold = True; r.font.color.rgb = AMBER_T; _ea(r)

# 反応行：起こる工程の真下に配置（どこで・どれが起こるか）
yr = 3.55; rh = 1.5; zbot = zy + zh
def conn(cx, color=GUIDE):
    rrect(s, cx-0.012, zbot, 0.024, yr-zbot, color)
# 上流：遊離アルデヒドは大半パージ・一部が別の形(前駆体)になって下流へ（EO反応系〜EG濃縮系 s0〜s2）
fl = lefts[0]; fr = lefts[2] + bw
vbox(s, fl, yr, fr-fl, rh,
     [("アルデヒドが入る → 大半はパージで系外へ", 11, True, DBLUE),
      ("CH3CHO（原料EO＋触媒劣化で増加・軽い）", 10, False, BLACK),
      ("残りは別の形（色のもと＝前駆体）に変わって下流へ", 10, False, INK),
      ("→ 濃縮塔の底では遊離アルデヒドとして検出されない", 9.5, True, AMBER_T)],
     H_FILL, line=BLUE)
conn((fl+fr)/2)
def rbox(ci, w, lines, fill=H_FILL, line=BLUE, lw=1.1):
    l = max(0.3, min(cxs[ci] - w/2, W-0.3-w))
    vbox(s, l, yr, w, rh, lines, fill, line=line, lw=lw); conn(cxs[ci])
# EG脱水系（s3）：反応の始まり
rbox(3, 2.0, [("水が抜けて反応場に", 11.5, True, DBLUE), ("前駆体どうしが", 10, False, BLACK), ("つながって育ち始める", 10, False, BLACK)])
# MEG精製系（s4）：通過
vbox(s, cxs[4]-0.9, yr+0.35, 1.8, rh-0.7, [("通過", 11, True, GRAY), ("新たな反応なし", 9.5, False, GRAY)],
     RGBColor(0xF2,0xF4,0xF6), line=LINE, lw=0.9); conn(cxs[4], RGBColor(0xD9,0xDF,0xE4))
# DEG精製系（s5）：色のもと
rbox(5, 1.8, [("色のもと（前駆体）が育つ", 10, True, DBLUE), ("CH3-(CH=CH)3-CHO", 9.5, False, BLACK), ("まだ淡い・330nm", 9, False, GRAY)])
# 製品タンク（s6）：色相悪化物質
rbox(6, 1.7, [("色相悪化物質に育つ", 10.5, True, AMBER_T), ("CH3-(CH=CH)n-CHO", 9, False, BLACK), ("黄色・450nm", 10, True, AMBER_T)], fill=AMBER, line=AMBER_T, lw=1.25)

note = tbx(s, 0.4, 5.3, W-0.8, 1.7)
P(note, "・アルデヒドは軽いので濃縮系のパージで大半が系外へ抜ける。残りは別の形（色のもと＝前駆体）に変わって下流へ運ばれる＝濃縮塔の底では遊離アルデヒドとして検出されない（だから上流では色がつかない）。", 11.5, first=True, space=6)
P(note, "・水が少なくなる脱水塔以降で、その前駆体どうしがつながって育ち、製品タンクの長期保管でさらに育って黄色（450nm）になる。前駆体を使い切るとAPHA30くらいで頭打ち。", 11.5, space=6)
P(note, "・酸素があると（アルデヒドや前駆体が）有機酸に変わって反応が止まる（窒素タンクで進み、空気に触れると止まる）。MEG塔の温度を下げると分解されず重い形で下流（DEG）へ回る。", 11.5, space=6)

# このスライドをメカニズム(S3)の直後へ移動
sldIdLst = prs.slides._sldIdLst
ids = list(sldIdLst)
newid = ids[-1]
sldIdLst.remove(newid)
sldIdLst.insert(3, newid)  # 0-based: 表紙(0)発生事象(1)メカニズム(2)の後＝index3

# ── 巻末：補足資料一覧／用意・確認する資料 ──
def add_table_slide(title, intro, headers, rows, widths):
    s = prs.slides.add_slide(ly)
    for ph in list(s.placeholders):
        if ph.placeholder_format.idx != 0 and ph.placeholder_format.type in (
            PP_PLACEHOLDER.BODY, PP_PLACEHOLDER.OBJECT, PP_PLACEHOLDER.SUBTITLE, PP_PLACEHOLDER.CENTER_TITLE):
            ph._element.getparent().remove(ph._element)
    for ph in s.placeholders:
        if ph.placeholder_format.idx == 0:
            ph.text = ""; r = ph.text_frame.paragraphs[0].add_run(); r.text = title
            r.font.size = Pt(24); r.font.bold = True; r.font.color.rgb = BLACK; _ea(r)
    if intro:
        P(tbx(s, 0.5, 1.35, W-1.0, 0.5), intro, 11, color=GRAY, first=True)
    nr = len(rows)+1; nc = len(headers)
    tb = s.shapes.add_table(nr, nc, Emu(int(0.5*IN)), Emu(int(1.9*IN)), Emu(int((W-1.0)*IN)), Emu(int(4.8*IN))).table
    for j, wd in enumerate(widths): tb.columns[j].width = Emu(int(wd*IN))
    def cell(i, j, text, bold=False, fill=WHITE, color=INK, size=10.5):
        c = tb.cell(i, j); c.text = ""; c.vertical_anchor = MSO_ANCHOR.MIDDLE
        c.margin_left = Emu(int(0.05*IN)); c.margin_top = Emu(int(0.02*IN)); c.margin_bottom = Emu(int(0.02*IN))
        p = c.text_frame.paragraphs[0]; r = p.add_run(); r.text = text
        r.font.size = Pt(size); r.font.bold = bold; r.font.color.rgb = color; _ea(r)
        c.fill.solid(); c.fill.fore_color.rgb = fill
    for j, h in enumerate(headers): cell(0, j, h, bold=True, color=DBLUE, fill=H_FILL, size=11)
    for i, row in enumerate(rows, 1):
        for j, v in enumerate(row): cell(i, j, v, fill=(WHITE if i % 2 else RGBColor(0xF7,0xFA,0xFD)), size=10)
    return s

add_table_slide(
    "（巻末）補足資料の一覧　※本編の流れで使うもの・聞かれた時に出すもの",
    "本編：発生事象→推定メカニズム→工程ごとの事実（運転・設備・基礎研）→計画停止対応→S/U時の対応。下記は補足として後ろに置く。",
    ["補足資料（タイトル）", "記載する内容", "状態"],
    [
        ["検査箇所・機器開放要否と結果", "EG精製系の開放箇所と、開けなくてよい理由（脱水塔以降が縮合の場）。C-1503以外に異常なし。", "あり"],
        ["サンプル分析（鉄錆と450nm）", "付着物・タンク異物はほぼ鉄サビ。DEG添加で380nmのみ＝450nmを作らない。", "あり"],
        ["運転対応実績とDEG色相への影響", "NaOH・温度・回収率・還流などの影響（『改善効果』でなく『影響』表現）。温度↓・NaOHはS/U条件の根拠。", "要整理"],
        ["MEG塔リボイラー(E-1501)穴と水分", "2025/10〜の水混入で水分が高めに見えた経緯。修理後は低めで推移。※本筋を混乱させるため聞かれた時のみ。", "あり"],
        ["AALD/FALDバランス", "C-1502で見かけ上ALDが生成（塔間循環）。", "あり"],
        ["UV330nm 活性剤投入有無の比較", "早期色相判断の可否検討。停止前（活性剤なし）と5月（リン酸あり）の比較データ。", "要整理（製造）"],
        ["リン酸の到達位置（C-1503壁面P分析）", "リン酸が効いた場所の絞り込み。XRFでP検出有無。", "結果待ち（基礎研）"],
        ["DEG/TEGリサイクル系のALDバランス", "上流から下流まで関係箇所を網羅したことを示す。", "分析待ち"],
    ],
    [4.2, 6.6, 1.5],
)

add_table_slide(
    "対策会議までに用意・確認する資料（佐野さんと整理）",
    "メカニズム→開放の紐付けを月曜にRD・木村含めて詰める。下記を揃える。",
    ["項目", "内容・狙い", "担当・期限"],
    [
        ["メカニズムの確定", "推定メカニズムに基づき開放箇所を紐づける筋を固める（金属触媒でなくアルデヒド由来）。", "月曜：佐野・高橋・RD・木村"],
        ["リン酸の到達位置", "C-1503壁面のP分析（XRF）。反応場の絞り込み材料。", "基礎研（野田）・近日"],
        ["DEG/TEGリサイクル系バランス", "関係する上流〜下流を網羅したと示すALDバランス。", "S/U会議まで・要レビュー"],
        ["UV330nmの比較データ", "活性剤投入有無での早期判断の可否。比較対象（トラブル前データ）も。", "製造・データ整理"],
        ["初期流動品の監視運用", "一発目はタンクで1週間〜10日UV経時を確認。運転条件変更時の扱いも。", "品証と来週調整"],
        ["運転対応実績の整理", "温度・NaOH・回収率・還流の『影響』。改善効果ありと書かない。", "本編から補足へ"],
    ],
    [3.6, 6.4, 2.3],
)

prs.save(OUT)
import zipfile, collections
names = zipfile.ZipFile(OUT).namelist()
dup = [k for k, v in collections.Counter(names).items() if v > 1]
print("saved:", os.path.normpath(OUT), "| slides:", len(prs.slides._sldIdLst), "| dup:", dup or "なし")
print("S6 filled:", bool(s6), "| S11 filled:", bool(s11))
